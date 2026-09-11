from datetime import datetime
from io import BytesIO
from textwrap import fill
from xml.sax.saxutils import escape

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import pandas as pd
from matplotlib.patches import Patch
from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from config import EXPORT_DIR


BLUE_HEX = "#345998"
LIGHT_BLUE_HEX = "#EAF3FC"
MONITORING_TABLE_BLUE_HEX = "#B4D2F0"
PAGE_BACKGROUND_HEX = "#F3F7FC"
STATUS_COLORS = {
    "Solucionado": "#4F81BD",
    "Repactuado": "#FFF200",
    "A vencer": "#00B050",
    "Vencido": "#FF0000",
    "Sem Plano": "#BFBFBF",
}
ORIGIN_COLORS = {
    "Auditoria Interna": "#7B337E",
    "Auditoria Externa": "#CA043E",
    "Auditoria Independente": "#FD7600",
    "TCU/CGU": "#FED000",
}
REPACTUATION_COLORS = {
    "1": "#92D050",
    "2": "#FFF200",
    "3": "#FFC000",
    "Acima 3": "#FF0000",
}
ACTIVE_PLAN_COLORS = {
    "Planos repactuados": "#F5F506",
    "Demais planos ativos": "#1565C0",
}
KPI_COLORS = ["#FFC000", "#5B9BD5", "#FF0000"]
SEMAFORO_CATEGORY_COLORS = {
    "Atenção": "#FF3333",
    "Monitoramento": "#FFDE59",
    "Conhecimento": "#76D84F",
}
REINCIDENCE_KPI_COLORS = {
    "Planos em Monitoramento": "#5B9BD5",
    "Planos Reincidentes": "#FF3333",
    "Taxa de Reincidência": "#FFDE59",
}
REINCIDENCE_RISK_COLORS = {
    "Alto": "#FF3333",
    "Significativo": "#F59A2F",
    "Moderado": "#FFDE59",
    "Baixo": "#76D84F",
}
SEMAFORO_RISK_COLORS = {
    "Alto": "#FF3333",
    "Significativo": "#F59A2F",
    "Moderado": "#F8FF25",
    "Baixo": "#76D84F",
}
REINCIDENCE_REPORT_COLORS = [
    "#FDC01D",
    "#FF8A1A",
    "#F23B1D",
    "#A7265D",
]
PAINT_STATUS_COLORS = {
    "Concluída": "#76D84F",
    "Em Andamento": "#FFDE59",
    "Não Iniciada": "#FF3333",
}
PAINT_KPI_COLORS = ["#5B9BD5", "#76D84F", "#FFDE59", "#FF3333"]
CARTEIRA_KPI_COLORS = ["#A7B0BD", "#5B9BD5", "#76D84F", "#FF3333"]
CARTEIRA_STATUS_COLORS = {
    "Auditorias com Planos em Monitoramento": "#5B9BD5",
    "Auditorias Encerradas": "#76D84F",
    "Aguardando Plano de Ação": "#FF3333",
}
CARTEIRA_YEAR_COLORS = [
    "#6C63E8",
    "#FF5C7A",
    "#27C2A4",
    "#FF9F1C",
    "#2E86DE",
    "#C846D9",
]

BLUE = colors.HexColor(BLUE_HEX)
LIGHT_BLUE = colors.HexColor(LIGHT_BLUE_HEX)
MONITORING_TABLE_BLUE = colors.HexColor(MONITORING_TABLE_BLUE_HEX)
THIN_GREY = colors.HexColor("#9BAAC2")
def _timestamp():
    return datetime.now().strftime("%Y%m%d_%H%M%S_%f")


def _filter_text(filters):
    return (
        "Nenhum filtro aplicado."
        if not filters
        else " | ".join(f"{key}: {value}" for key, value in filters.items())
    )


def _safe_text(value):
    if pd.isna(value):
        return ""
    if isinstance(value, pd.Timestamp):
        return value.strftime("%d/%m/%Y")
    return str(value)


def _text_color_for(background):
    clean = background.lstrip("#")
    red, green, blue = (int(clean[index:index + 2], 16) for index in (0, 2, 4))
    luminance = (0.299 * red) + (0.587 * green) + (0.114 * blue)
    return "#111111" if luminance >= 145 else "#FFFFFF"




def export_table_excel(df, filters, title=None, prefix="tabela_analitica"):
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    filename = EXPORT_DIR / f"{prefix}_{_timestamp()}.xlsx"

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Tabela Analítica", index=False, startrow=3)
        sheet = writer.book["Tabela Analítica"]
        resolved_title = title or (
            f"Tabela Analítica - {len(df)} Planos Ativos"
            if filters.get("Status dos planos") == "Planos ativos"
            else "Tabela Analítica de Planos de Ação"
        )
        last_column = get_column_letter(max(1, len(df.columns)))
        sheet.merge_cells(f"A1:{last_column}1")
        sheet["A1"] = resolved_title
        sheet["A1"].fill = PatternFill("solid", fgColor=BLUE_HEX.lstrip("#"))
        sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
        sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
        sheet["A2"] = _filter_text(filters)
        sheet["A2"].font = Font(size=9, italic=True, color="52647A")
        sheet.row_dimensions[1].height = 28
        sheet.freeze_panes = "A5"
        last_row = 4 + len(df)
        sheet.auto_filter.ref = f"A4:{last_column}{last_row}"
        sheet.sheet_view.showGridLines = False

        for cell in sheet[4]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=BLUE_HEX.lstrip("#"))
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for column_index, column_cells in enumerate(sheet.iter_cols(), start=1):
            maximum = max(len(_safe_text(cell.value)) for cell in column_cells)
            sheet.column_dimensions[get_column_letter(column_index)].width = max(
                12, min(45, maximum + 2)
            )

        for row in range(5, last_row + 1):
            sheet.row_dimensions[row].height = 36
            for cell in sheet[row]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_title_rows = "1:4"
        sheet.print_area = f"A1:{last_column}{last_row}"

    return filename


def _base_doc(filename):
    return SimpleDocTemplate(
        str(filename),
        pagesize=landscape(A3),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
        title="Indicadores de Monitoramento",
        author="Dashboard de Monitoramento",
    )


def _title_bar(title):
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "AnalyticalTitleBar",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=17,
        leading=20,
        textColor=colors.white,
        alignment=1,
    )
    available = landscape(A3)[0] - (20 * mm)
    bar = Table(
        [[Paragraph(escape(title), title_style)]],
        colWidths=[available],
        rowHeights=[12 * mm],
    )
    bar.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), BLUE),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return bar


def _title(elements, title, filters):
    styles = getSampleStyleSheet()
    elements.extend(
        [
            _title_bar(title),
            Spacer(1, 3 * mm),
            Paragraph(escape(_filter_text(filters)), styles["Normal"]),
            Spacer(1, 8),
        ]
    )


def _styled_table(data, col_widths=None, font_size=7):
    table = Table(data, repeatRows=1, colWidths=col_widths, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), BLUE),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), font_size),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.25, THIN_GREY),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT_BLUE]),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    return table


def export_table_pdf(df, filters, title=None, prefix="tabela_analitica"):
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    filename = EXPORT_DIR / f"{prefix}_{_timestamp()}.pdf"
    doc = _base_doc(filename)
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle(
        "SmallCell",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=5.5,
        leading=6.5,
    )
    elements = []
    resolved_title = title or (
        f"Tabela Analítica - {len(df)} Planos Ativos"
        if filters.get("Status dos planos") == "Planos ativos"
        else "Tabela Analítica de Planos de Ação"
    )
    _title(elements, resolved_title, filters)
    columns = df.columns.tolist()
    chunk_size = 6

    for start in range(0, len(columns), chunk_size):
        chunk = columns[start:start + chunk_size]
        elements.append(
            Paragraph(
                f"Colunas {start + 1} a {start + len(chunk)} de {len(columns)}",
                styles["Heading2"],
            )
        )
        data = [
            [Paragraph("Registro", cell_style)]
            + [Paragraph(escape(str(column)), cell_style) for column in chunk]
        ]
        for index, row in df[chunk].iterrows():
            data.append(
                [Paragraph(str(index + 1), cell_style)]
                + [
                    Paragraph(
                        escape(_safe_text(row[column])).replace("\n", "<br/>"),
                        cell_style,
                    )
                    for column in chunk
                ]
            )

        available = landscape(A3)[0] - (20 * mm)
        elements.append(
            _styled_table(
                data,
                [18 * mm] + [(available - (18 * mm)) / len(chunk)] * len(chunk),
                5.5,
            )
        )
        if start + chunk_size < len(columns):
            elements.append(PageBreak())
            _title(elements, resolved_title, filters)

    doc.build(elements, onFirstPage=_pdf_page_footer, onLaterPages=_pdf_page_footer)
    return filename




def _pdf_page_footer(canvas, document):
    canvas.saveState()
    page_width, _ = landscape(A3)
    canvas.setStrokeColor(colors.HexColor("#D8E2EF"))
    canvas.line(10 * mm, 7 * mm, page_width - (10 * mm), 7 * mm)
    canvas.setFillColor(colors.HexColor("#52647A"))
    canvas.setFont("Helvetica", 7)
    canvas.drawString(10 * mm, 4 * mm, "Dashboard de Monitoramento")
    canvas.drawRightString(
        page_width - (10 * mm),
        4 * mm,
        f"Página {document.page}",
    )
    canvas.restoreState()


def _pdf_header(section, filters, title="Indicadores de Monitoramento"):
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "DashboardTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=21,
        textColor=colors.HexColor("#0E2146"),
        spaceAfter=3,
    )
    section_style = ParagraphStyle(
        "DashboardSection",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=12,
        textColor=BLUE,
        spaceAfter=2,
    )
    filter_style = ParagraphStyle(
        "DashboardFilters",
        parent=styles["BodyText"],
        fontSize=7.5,
        leading=9,
        textColor=colors.HexColor("#52647A"),
    )
    return [
        Paragraph(escape(title), title_style),
        Paragraph(escape(section), section_style),
        Paragraph(escape(_filter_text(filters)), filter_style),
        Spacer(1, 5 * mm),
    ]


def _paragraph(text, size=8, bold=False, color="#111111", align=0):
    styles = getSampleStyleSheet()
    return Paragraph(
        escape(str(text)),
        ParagraphStyle(
            f"Cell_{size}_{bold}_{color}_{align}",
            parent=styles["BodyText"],
            fontName="Helvetica-Bold" if bold else "Helvetica",
            fontSize=size,
            leading=size + 2,
            textColor=colors.HexColor(color),
            alignment=align,
        ),
    )


def _kpi_cards_pdf(indicators):
    cards = [
        ("TOTAL DE PLANOS DE AÇÃO", indicators.get("total_planos_acao", 0)),
        ("PLANOS SOLUCIONADOS", indicators.get("planos_solucionados", 0)),
        ("PLANOS VENCIDOS", indicators.get("planos_vencidos", 0)),
    ]
    cell_contents = []
    for label, value in cards:
        cell_contents.append(
            [
                _paragraph(label, size=10, bold=True, align=1),
                Spacer(1, 4),
                _paragraph(value, size=25, bold=True, align=1),
            ]
        )

    table = Table(
        [cell_contents],
        colWidths=[128 * mm] * 3,
        rowHeights=[30 * mm],
        hAlign="CENTER",
    )
    commands = [
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("LEFTPADDING", (0, 0), (-1, -1), 12),
        ("RIGHTPADDING", (0, 0), (-1, -1), 12),
        ("TOPPADDING", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#B8C7D9")),
        ("INNERGRID", (0, 0), (-1, -1), 6, colors.white),
        ("ROUNDEDCORNERS", [9]),
    ]
    for column, color in enumerate(KPI_COLORS):
        commands.append(("BACKGROUND", (column, 0), (column, 0), colors.HexColor(color)))
    table.setStyle(TableStyle(commands))
    return table


def _pie_chart_image(
    labels,
    values,
    title,
    palette,
    doughnut=False,
    show_value_with_percent=False,
):
    fig, axis = plt.subplots(figsize=(6.8, 4.2))
    color_list = [palette.get(label, BLUE_HEX) for label in labels]
    total = sum(values)

    def percent_label(percent):
        if percent <= 0 or not total:
            return ""
        percentage = f"{percent:.1f}".replace(".", ",")
        if show_value_with_percent:
            value = int(round(percent * total / 100))
            return f"{value} ({percentage}%)"
        return f"{percentage}%"

    wedge = {"edgecolor": "white", "linewidth": 1.5}
    if doughnut:
        wedge["width"] = 0.46

    axis.pie(
        values,
        colors=color_list,
        startangle=90,
        autopct=percent_label,
        pctdistance=0.72,
        textprops={"fontsize": 8, "fontweight": "bold", "color": "#111111"},
        wedgeprops=wedge,
    )
    axis.axis("equal")
    axis.set_title(title, fontsize=11, fontweight="bold", color="#111111", pad=8)
    axis.legend(
        labels,
        loc="lower center",
        bbox_to_anchor=(0.5, -0.17),
        ncol=2,
        fontsize=7.2,
        frameon=True,
        facecolor="white",
        edgecolor="#D8E2EF",
    )
    fig.tight_layout()
    buffer = BytesIO()
    fig.savefig(buffer, format="png", dpi=170, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    buffer.seek(0)
    return buffer


def _stacked_status_chart_image(rows, group_column, title):
    labels = [str(row[group_column]) for row in rows]
    statuses = ["A vencer", "Repactuado", "Vencido", "Solucionado", "Sem Plano"]
    figure_height = max(5.4, min(10.2, (0.48 * max(len(labels), 1)) + 2.4))
    fig, axis = plt.subplots(figsize=(11.5, figure_height))
    left = [0] * len(labels)

    for status in statuses:
        values = [int(row.get(status, 0) or 0) for row in rows]
        bars = axis.barh(
            labels,
            values,
            left=left,
            label=status,
            color=STATUS_COLORS[status],
            edgecolor="white",
            linewidth=0.8,
        )
        for bar, value in zip(bars, values):
            if value > 0:
                axis.text(
                    bar.get_x() + (bar.get_width() / 2),
                    bar.get_y() + (bar.get_height() / 2),
                    str(value),
                    ha="center",
                    va="center",
                    fontsize=8,
                    fontweight="bold",
                    color="#111111",
                )
        left = [old + new for old, new in zip(left, values)]

    axis.set_xlabel("Quantidade de planos", fontsize=9)
    axis.set_title(title, fontsize=11, fontweight="bold")
    axis.grid(axis="x", alpha=0.18)
    if group_column == "Diretoria":
        axis.tick_params(axis="x", which="both", labelbottom=False)
    axis.legend(
        loc="lower center",
        bbox_to_anchor=(0.5, -0.22),
        ncol=5,
        fontsize=7.5,
        frameon=True,
        facecolor="white",
        edgecolor="#D8E2EF",
    )
    axis.invert_yaxis()
    fig.tight_layout()
    buffer = BytesIO()
    fig.savefig(buffer, format="png", dpi=170, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    buffer.seek(0)
    return buffer


def _period_chart_image(rows):
    labels = [str(row.get("periodo", "")) for row in rows]
    values = [int(row.get("quantidade", 0) or 0) for row in rows]
    figure_height = max(5.4, min(10.2, (0.42 * max(len(labels), 1)) + 2.2))
    fig, axis = plt.subplots(figsize=(10.5, figure_height))
    axis.barh(
        labels,
        values,
        color="#4F81BD",
        edgecolor="#345998",
        linewidth=1.0,
    )
    axis.set_xlabel("Quantidade de planos", fontsize=9)
    axis.set_title(
        "Gráfico do Período Previsto de Conclusão de Plano",
        fontsize=11,
        fontweight="bold",
    )
    axis.grid(axis="x", alpha=0.18)
    axis.invert_yaxis()

    maximum = max(values, default=0)
    axis.set_xlim(0, max(1, maximum + 1))
    fig.tight_layout()
    buffer = BytesIO()
    fig.savefig(buffer, format="png", dpi=170, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    buffer.seek(0)
    return buffer


def _overview_pdf_table(mapping, first_header, palette):
    labels = list(mapping.keys())
    values = [int(mapping[label].get("quantidade", 0) or 0) for label in labels]
    data = [[first_header, "Quant."]]
    data.extend([[label, value] for label, value in zip(labels, values)])
    data.append(["TOTAL", sum(values)])
    table = Table(data, colWidths=[48 * mm, 18 * mm], rowHeights=8.2 * mm)
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), MONITORING_TABLE_BLUE),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#111111")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("BACKGROUND", (0, -1), (-1, -1), BLUE),
        ("TEXTCOLOR", (0, -1), (-1, -1), colors.HexColor("#111111")),
        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("GRID", (0, 0), (-1, -1), 0.45, colors.HexColor("#737B85")),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ]
    for index, label in enumerate(labels, start=1):
        background = palette.get(label, LIGHT_BLUE_HEX)
        commands.extend(
            [
                ("BACKGROUND", (0, index), (0, index), colors.HexColor(background)),
                (
                    "TEXTCOLOR",
                    (0, index),
                    (0, index),
                    colors.HexColor("#111111"),
                ),
                ("FONTNAME", (0, index), (0, index), "Helvetica-Bold"),
            ]
        )
    table.setStyle(TableStyle(commands))
    return table


def _group_export_table(rows, group_column):
    headers = [
        group_column,
        "A vencer",
        "Repactuado",
        "Vencido",
        "Solucionado",
        "Sem Plano",
        "Total",
        "%",
    ]
    data = [headers]
    for row in rows:
        data.append(
            [
                row[group_column],
                row["A vencer"],
                row["Repactuado"],
                row["Vencido"],
                row["Solucionado"],
                row["Sem Plano"],
                row["Total"],
                f'{float(row.get("%", 0)):.2f}%'.replace(".", ","),
            ]
        )
    total = sum(int(row.get("Total", 0)) for row in rows)
    data.append(
        [
            "TOTAL",
            sum(int(row.get("A vencer", 0)) for row in rows),
            sum(int(row.get("Repactuado", 0)) for row in rows),
            sum(int(row.get("Vencido", 0)) for row in rows),
            sum(int(row.get("Solucionado", 0)) for row in rows),
            sum(int(row.get("Sem Plano", 0)) for row in rows),
            total,
            "100,00%" if total else "0,00%",
        ]
    )
    return data


def _group_pdf_table(rows, group_column):
    data = _group_export_table(rows, group_column)
    widths = [48 * mm, 15 * mm, 17 * mm, 15 * mm, 17 * mm, 17 * mm, 14 * mm, 14 * mm]
    table = Table(data, colWidths=widths, repeatRows=1)
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), BLUE),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#111111")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("BACKGROUND", (0, -1), (-1, -1), BLUE),
        ("TEXTCOLOR", (0, -1), (-1, -1), colors.HexColor("#111111")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, LIGHT_BLUE]),
        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTSIZE", (0, 0), (-1, -1), 6.4),
        ("GRID", (0, 0), (-1, -1), 0.35, THIN_GREY),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    for column, status in enumerate(
        ["A vencer", "Repactuado", "Vencido", "Solucionado", "Sem Plano"],
        start=1,
    ):
        background = STATUS_COLORS[status]
        commands.extend(
            [
                ("BACKGROUND", (column, 0), (column, 0), colors.HexColor(background)),
                (
                    "TEXTCOLOR",
                    (column, 0),
                    (column, 0),
                    colors.HexColor("#111111"),
                ),
            ]
        )
    commands.extend(
        [
            ("BACKGROUND", (0, 0), (0, 0), MONITORING_TABLE_BLUE),
            ("BACKGROUND", (-2, 0), (-1, 0), MONITORING_TABLE_BLUE),
        ]
    )
    table.setStyle(TableStyle(commands))
    return table


def _repactuation_pdf_table(indicators):
    distribution = indicators.get("distribuicao_repactuacoes", {})
    rows = [
        ("1", int(distribution.get("1", 0) or 0)),
        ("2", int(distribution.get("2", 0) or 0)),
        ("3", int(distribution.get("3", 0) or 0)),
        ("Acima 3", int(distribution.get("Acima de 3", 0) or 0)),
    ]
    total = sum(value for _, value in rows)
    data = [["Repactuações", "Quant. Planos", "%"]]
    data.extend(
        [
            label,
            value,
            f"{((value / total) * 100 if total else 0):.1f}".replace(".", ","),
        ]
        for label, value in rows
    )
    data.append(["TOTAL", total, "100,0" if total else "0,0"])
    table = Table(data, colWidths=[45 * mm, 38 * mm, 26 * mm], rowHeights=9 * mm)
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), MONITORING_TABLE_BLUE),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#111111")),
        ("BACKGROUND", (0, -1), (-1, -1), MONITORING_TABLE_BLUE),
        ("TEXTCOLOR", (0, -1), (-1, -1), colors.HexColor("#111111")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#5F6670")),
    ]
    for index, (label, _) in enumerate(rows, start=1):
        commands.append(
            (
                "BACKGROUND",
                (0, index),
                (-1, index),
                colors.HexColor(REPACTUATION_COLORS[label]),
            )
        )
        commands.append(("TEXTCOLOR", (0, index), (-1, index), colors.HexColor("#111111")))
    table.setStyle(TableStyle(commands))
    return table


def _period_pdf_table(rows):
    data = [["Planos Repactuados e Vencidos", "Quant. Planos"]]
    data.extend(
        [
            row.get("periodo", ""),
            int(row.get("quantidade", 0) or 0),
        ]
        for row in rows
    )
    table = Table(data, colWidths=[62 * mm, 47 * mm], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), MONITORING_TABLE_BLUE),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#111111")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.35, THIN_GREY),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return table


def _panel_pdf(title, content, width):
    title_flowable = _paragraph(title, size=11, bold=True, color="#111111", align=1)
    panel = Table([[title_flowable], [content]], colWidths=[width])
    panel.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                ("BOX", (0, 0), (-1, -1), 0.7, colors.HexColor("#D8E2EF")),
                ("ROUNDEDCORNERS", [10]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, 0), 8),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 7),
                ("TOPPADDING", (0, 1), (-1, 1), 5),
                ("BOTTOMPADDING", (0, 1), (-1, 1), 8),
            ]
        )
    )
    return panel


def _repactuation_kpi_cards_pdf(indicators):
    cards = [
        ("PLANOS REPACTUADOS", int(indicators.get("planos_repactuados", 0) or 0)),
        (
            "TAXA DE REPACTUAÇÃO",
            f'{float(indicators.get("taxa_repactuacao", 0) or 0):.2f}%'.replace(".", ","),
        ),
    ]
    content = [
        [
            _paragraph(label, size=11, bold=True, align=1),
            Spacer(1, 4),
            _paragraph(value, size=24, bold=True, align=1),
        ]
        for label, value in cards
    ]
    inner = Table(
        [content],
        colWidths=[130 * mm, 130 * mm],
        rowHeights=[32 * mm],
        hAlign="CENTER",
    )
    inner.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FFF200")),
                ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#C8B900")),
                ("INNERGRID", (0, 0), (-1, -1), 12, colors.white),
                ("ROUNDEDCORNERS", [9]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("LEFTPADDING", (0, 0), (-1, -1), 14),
                ("RIGHTPADDING", (0, 0), (-1, -1), 14),
                ("TOPPADDING", (0, 0), (-1, -1), 10),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
            ]
        )
    )
    outer = Table([[inner]], colWidths=[390 * mm])
    outer.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                ("BOX", (0, 0), (-1, -1), 0.7, colors.HexColor("#D8E2EF")),
                ("ROUNDEDCORNERS", [10]),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("TOPPADDING", (0, 0), (-1, -1), 10),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
            ]
        )
    )
    return outer


def export_indicators_pdf(indicators, filters):
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    filename = EXPORT_DIR / f"indicadores_{_timestamp()}.pdf"
    document = _base_doc(filename)
    elements = []

    elements.extend(_pdf_header("Resumo executivo e visões gerais", filters))
    elements.extend([_kpi_cards_pdf(indicators), Spacer(1, 6 * mm)])

    status = indicators.get("visao_geral_status", {})
    status_labels = list(status.keys())
    status_values = [int(status[label].get("quantidade", 0) or 0) for label in status_labels]
    origin = indicators.get("visao_geral_origem", {})
    origin_labels = list(origin.keys())
    origin_values = [int(origin[label].get("quantidade", 0) or 0) for label in origin_labels]

    status_content = Table(
        [
            [
                _overview_pdf_table(status, "Situação", STATUS_COLORS),
                Image(
                    _pie_chart_image(
                        status_labels,
                        status_values,
                        "Visão Geral dos Planos de Ação",
                        STATUS_COLORS,
                    ),
                    width=118 * mm,
                    height=72 * mm,
                ),
            ]
        ],
        colWidths=[67 * mm, 120 * mm],
    )
    origin_content = Table(
        [
            [
                _overview_pdf_table(origin, "Origem", ORIGIN_COLORS),
                Image(
                    _pie_chart_image(
                        origin_labels,
                        origin_values,
                        "Visão Geral de Origem",
                        ORIGIN_COLORS,
                    ),
                    width=118 * mm,
                    height=72 * mm,
                ),
            ]
        ],
        colWidths=[67 * mm, 120 * mm],
    )
    overviews = Table(
        [
            [
                _panel_pdf("Visão Geral dos Planos de Ação", status_content, 194 * mm),
                _panel_pdf("Visão Geral de Origem", origin_content, 194 * mm),
            ]
        ],
        colWidths=[196 * mm, 196 * mm],
    )
    overviews.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 2),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    elements.append(overviews)

    directorate_rows = indicators.get("planos_por_diretoria", [])
    elements.append(PageBreak())
    elements.extend(_pdf_header("Planos de Ação por Diretoria", filters))
    directorate_content = Table(
        [
            [
                _group_pdf_table(directorate_rows, "Diretoria"),
                Image(
                    _stacked_status_chart_image(
                        directorate_rows,
                        "Diretoria",
                        "Gráfico de Planos por Diretoria",
                    ),
                    width=222 * mm,
                    height=147 * mm,
                ),
            ]
        ],
        colWidths=[158 * mm, 224 * mm],
    )
    elements.append(
        _panel_pdf("Tabela e gráfico por Diretoria", directorate_content, 390 * mm)
    )

    unit_rows = indicators.get("planos_por_unidade", [])
    elements.append(PageBreak())
    elements.extend(_pdf_header("Planos de Ação por Unidade/Agência", filters))
    unit_content = Table(
        [
            [
                _group_pdf_table(unit_rows, "Unidade/Agência"),
                Image(
                    _stacked_status_chart_image(
                        unit_rows,
                        "Unidade/Agência",
                        "Gráfico de Planos por Unidade/Agência",
                    ),
                    width=222 * mm,
                    height=147 * mm,
                ),
            ]
        ],
        colWidths=[158 * mm, 224 * mm],
    )
    elements.append(
        _panel_pdf("Tabela e gráfico por Unidade/Agência", unit_content, 390 * mm)
    )

    elements.append(PageBreak())
    elements.extend(_pdf_header("Repactuações e Planos Ativos", filters))
    active_labels = ["Planos repactuados", "Demais planos ativos"]
    active_values = [
        int(indicators.get("planos_repactuados", 0) or 0),
        max(
            0,
            int(indicators.get("planos_ativos", 0) or 0)
            - int(indicators.get("planos_repactuados", 0) or 0),
        ),
    ]
    repactuation_content = Table(
        [
            [
                _panel_pdf(
                    "Repactuações e Planos Ativos",
                    _repactuation_pdf_table(indicators),
                    194 * mm,
                ),
                _panel_pdf(
                    "Gráfico de Repactuações e Planos Ativos",
                    Image(
                        _pie_chart_image(
                            active_labels,
                            active_values,
                            "Gráfico de Repactuações e Planos Ativos",
                            ACTIVE_PLAN_COLORS,
                            doughnut=True,
                            show_value_with_percent=True,
                        ),
                        width=178 * mm,
                        height=93 * mm,
                    ),
                    194 * mm,
                ),
            ]
        ],
        colWidths=[196 * mm, 196 * mm],
    )
    repactuation_content.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 2),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    elements.extend(
        [
            _repactuation_kpi_cards_pdf(indicators),
            Spacer(1, 7 * mm),
            repactuation_content,
        ]
    )

    period_rows = indicators.get(
        "periodos_conclusao_planos",
        indicators.get("repactuacoes_por_periodo", []),
    )
    elements.append(PageBreak())
    elements.extend(_pdf_header("Período Previsto para Conclusão de Planos", filters))
    period_content = Table(
        [
            [
                _period_pdf_table(period_rows),
                Image(
                    _period_chart_image(period_rows),
                    width=225 * mm,
                    height=147 * mm,
                ),
            ]
        ],
        colWidths=[112 * mm, 230 * mm],
    )
    elements.append(
        _panel_pdf("Tabela e gráfico por período", period_content, 390 * mm)
    )

    document.build(
        elements,
        onFirstPage=_pdf_page_footer,
        onLaterPages=_pdf_page_footer,
    )
    return filename




def _excel_fill(color):
    return PatternFill("solid", fgColor=color.lstrip("#"))


def _excel_font_color(background):
    return _text_color_for(background).lstrip("#")


def _paint_range(sheet, min_row, max_row, min_col, max_col, fill):
    for row in sheet.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    ):
        for cell in row:
            cell.fill = fill


def _style_panel_range(sheet, min_row, max_row, min_col, max_col):
    white_fill = _excel_fill("#FFFFFF")
    border_color = "D8E2EF"
    thin = Side(style="thin", color=border_color)
    _paint_range(sheet, min_row, max_row, min_col, max_col, white_fill)
    for row in range(min_row, max_row + 1):
        for column in range(min_col, max_col + 1):
            cell = sheet.cell(row, column)
            cell.border = Border(
                left=thin if column == min_col else cell.border.left,
                right=thin if column == max_col else cell.border.right,
                top=thin if row == min_row else cell.border.top,
                bottom=thin if row == max_row else cell.border.bottom,
            )


def _merge_title(sheet, row, title, first_col=1, last_col=19):
    _paint_range(sheet, row, row + 1, first_col, last_col, _excel_fill("#FFFFFF"))
    sheet.merge_cells(
        start_row=row,
        start_column=first_col,
        end_row=row + 1,
        end_column=last_col,
    )
    cell = sheet.cell(row, first_col, title)
    cell.font = Font(size=16, bold=True, color="0E2146")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[row].height = 22
    sheet.row_dimensions[row + 1].height = 12


def _section_title(sheet, row, title, first_col=1, last_col=19):
    _paint_range(sheet, row, row, first_col, last_col, _excel_fill(BLUE_HEX))
    sheet.merge_cells(
        start_row=row,
        start_column=first_col,
        end_row=row,
        end_column=last_col,
    )
    cell = sheet.cell(row, first_col, title)
    cell.font = Font(size=13, bold=True, color="FFFFFF")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[row].height = 24


def _excel_kpi_card(sheet, min_row, max_row, min_col, max_col, label, value, color):
    _paint_range(sheet, min_row, max_row, min_col, max_col, _excel_fill(color))
    sheet.merge_cells(
        start_row=min_row,
        start_column=min_col,
        end_row=max_row,
        end_column=max_col,
    )
    cell = sheet.cell(min_row, min_col)
    cell.value = f"{label}\n{value}"
    cell.font = Font(size=16, bold=True, color=_excel_font_color(color))
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="B8C7D9")
    for row in range(min_row, max_row + 1):
        for column in range(min_col, max_col + 1):
            sheet.cell(row, column).border = Border(
                left=thin if column == min_col else Side(style=None),
                right=thin if column == max_col else Side(style=None),
                top=thin if row == min_row else Side(style=None),
                bottom=thin if row == max_row else Side(style=None),
            )


def _write_excel_table(
    sheet,
    start_row,
    start_col,
    data,
    *,
    category_colors=None,
    category_column=0,
    status_header=False,
    repactuation_rows=False,
    period_rows=False,
    light_header=False,
):
    thin = Side(style="thin", color="9BAAC2")
    base_header_fill = _excel_fill(
        MONITORING_TABLE_BLUE_HEX
        if repactuation_rows or period_rows
        else BLUE_HEX
    )
    header_fill = (
        _excel_fill(MONITORING_TABLE_BLUE_HEX)
        if light_header
        else base_header_fill
    )

    for row_offset, values in enumerate(data):
        for column_offset, value in enumerate(values):
            cell = sheet.cell(start_row + row_offset, start_col + column_offset, value)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(
                horizontal="left" if column_offset == 0 else "center",
                vertical="center",
                wrap_text=True,
            )
            cell.font = Font(size=9, color="111111")

    column_count = len(data[0])
    for column_offset in range(column_count):
        cell = sheet.cell(start_row, start_col + column_offset)
        cell.fill = header_fill
        cell.font = Font(size=9, bold=True, color="111111")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    if status_header:
        statuses = ["A vencer", "Repactuado", "Vencido", "Solucionado", "Sem Plano"]
        for index, status in enumerate(statuses, start=1):
            cell = sheet.cell(start_row, start_col + index)
            color = STATUS_COLORS[status]
            cell.fill = _excel_fill(color)
            cell.font = Font(size=9, bold=True, color="111111")

    if category_colors:
        for row_offset, values in enumerate(data[1:-1], start=1):
            label = str(values[category_column])
            color = category_colors.get(label)
            if color:
                cell = sheet.cell(start_row + row_offset, start_col + category_column)
                cell.fill = _excel_fill(color)
                cell.font = Font(size=9, bold=True, color="111111")

    if repactuation_rows:
        for row_offset, values in enumerate(data[1:-1], start=1):
            color = REPACTUATION_COLORS.get(str(values[0]))
            if color:
                for column_offset in range(column_count):
                    cell = sheet.cell(start_row + row_offset, start_col + column_offset)
                    cell.fill = _excel_fill(color)
                    cell.font = Font(size=9, bold=True, color="111111")
        for column_offset in range(column_count):
            header = sheet.cell(start_row, start_col + column_offset)
            header.font = Font(size=9, bold=True, color="111111")

    if period_rows:
        for column_offset in range(column_count):
            sheet.cell(start_row, start_col + column_offset).font = Font(
                size=9, bold=True, color="111111"
            )
        for row_offset in range(1, len(data)):
            if row_offset % 2 == 0:
                for column_offset in range(column_count):
                    sheet.cell(
                        start_row + row_offset,
                        start_col + column_offset,
                    ).fill = _excel_fill("#F2F2F2")

    if len(data) > 1 and str(data[-1][0]).upper() == "TOTAL":
        for column_offset in range(column_count):
            cell = sheet.cell(start_row + len(data) - 1, start_col + column_offset)
            cell.fill = base_header_fill
            cell.font = Font(size=9, bold=True, color="111111")

    return start_row + len(data) - 1


def _chart_data_block(helper, start_row, title, headers, rows):
    helper.cell(start_row, 1, title)
    for column, header in enumerate(headers, start=1):
        helper.cell(start_row + 1, column, header)
    for row_offset, values in enumerate(rows, start=2):
        for column, value in enumerate(values, start=1):
            helper.cell(start_row + row_offset, column, value)
    return start_row + 1, start_row + 1 + len(rows)


def _apply_point_colors(chart, palette):
    if not chart.series:
        return
    points = []
    for index, color in enumerate(palette):
        point = DataPoint(idx=index)
        point.graphicalProperties.solidFill = color.lstrip("#")
        point.graphicalProperties.line.solidFill = "FFFFFF"
        points.append(point)
    chart.series[0].dPt = points


def _apply_series_colors(chart, palette):
    for series, color in zip(chart.series, palette):
        clean = color.lstrip("#")
        series.graphicalProperties.solidFill = clean
        series.graphicalProperties.line.solidFill = clean


def _add_pie_chart(
    sheet,
    helper,
    helper_start,
    anchor,
    title,
    labels,
    values,
    palette,
    *,
    doughnut=False,
    width=13.2,
    height=7.8,
):
    chart_image = XLImage(
        _pie_chart_image(
            labels,
            values,
            title,
            dict(zip(labels, palette)),
            doughnut=doughnut,
            show_value_with_percent=doughnut,
        )
    )
    chart_image.width = int(width * 37.8)
    chart_image.height = int(height * 37.8)
    sheet.add_image(chart_image, anchor)
    return helper_start + 1


def _add_group_chart(
    sheet,
    helper,
    helper_start,
    anchor,
    title,
    rows,
    group_column,
    *,
    width=21.5,
    height=10.5,
):
    chart_image = XLImage(
        _stacked_status_chart_image(
            rows,
            group_column,
            title,
        )
    )
    chart_image.width = int(width * 37.8)
    chart_image.height = int(height * 37.8)
    sheet.add_image(chart_image, anchor)
    return helper_start + 1


def _add_period_chart(
    sheet,
    helper,
    helper_start,
    anchor,
    rows,
    *,
    width=21.5,
    height=10.5,
):
    chart_image = XLImage(_period_chart_image(rows))
    chart_image.width = int(width * 37.8)
    chart_image.height = int(height * 37.8)
    sheet.add_image(chart_image, anchor)
    return helper_start + 1


def export_indicators_excel(indicators, filters):
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    filename = EXPORT_DIR / f"indicadores_{_timestamp()}.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dashboard"
    helper = workbook.create_sheet("Dados dos gráficos")
    helper.sheet_state = "hidden"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"

    for column in range(1, 20):
        sheet.column_dimensions[get_column_letter(column)].width = (
            17 if column in {1, 11} else 11.5
        )
    for row in range(1, 180):
        sheet.row_dimensions[row].height = 20

    _paint_range(sheet, 1, 180, 1, 19, _excel_fill(PAGE_BACKGROUND_HEX))
    _merge_title(sheet, 1, "Indicadores de Monitoramento")
    sheet.merge_cells("A3:S3")
    sheet["A3"] = _filter_text(filters)
    sheet["A3"].font = Font(size=9, italic=True, color="52647A")
    sheet["A3"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    _excel_kpi_card(
        sheet,
        5,
        8,
        1,
        5,
        "TOTAL DE PLANOS DE AÇÃO",
        int(indicators.get("total_planos_acao", 0) or 0),
        KPI_COLORS[0],
    )
    _excel_kpi_card(
        sheet,
        5,
        8,
        8,
        12,
        "PLANOS SOLUCIONADOS",
        int(indicators.get("planos_solucionados", 0) or 0),
        KPI_COLORS[1],
    )
    _excel_kpi_card(
        sheet,
        5,
        8,
        15,
        19,
        "PLANOS VENCIDOS",
        int(indicators.get("planos_vencidos", 0) or 0),
        KPI_COLORS[2],
    )

    helper_row = 1
    print_break_rows = []

    section_row = 10
    _section_title(sheet, section_row, "Visões Gerais")
    _style_panel_range(sheet, section_row + 1, section_row + 16, 1, 9)
    _style_panel_range(sheet, section_row + 1, section_row + 16, 11, 19)
    sheet.merge_cells(
        start_row=section_row + 1,
        start_column=1,
        end_row=section_row + 1,
        end_column=9,
    )
    sheet.cell(section_row + 1, 1, "Visão Geral dos Planos de Ação")
    sheet.cell(section_row + 1, 1).font = Font(size=11, bold=True, color="111111")
    sheet.cell(section_row + 1, 1).alignment = Alignment(horizontal="center")
    sheet.merge_cells(
        start_row=section_row + 1,
        start_column=11,
        end_row=section_row + 1,
        end_column=19,
    )
    sheet.cell(section_row + 1, 11, "Visão Geral de Origem")
    sheet.cell(section_row + 1, 11).font = Font(size=11, bold=True, color="111111")
    sheet.cell(section_row + 1, 11).alignment = Alignment(horizontal="center")

    status = indicators.get("visao_geral_status", {})
    status_labels = list(status.keys())
    status_values = [int(status[label].get("quantidade", 0) or 0) for label in status_labels]
    status_data = [["Situação", "Quant."]]
    status_data.extend([[label, value] for label, value in zip(status_labels, status_values)])
    status_data.append(["TOTAL", sum(status_values)])
    _write_excel_table(
        sheet,
        section_row + 3,
        1,
        status_data,
        category_colors=STATUS_COLORS,
        light_header=True,
    )
    helper_row = _add_pie_chart(
        sheet,
        helper,
        helper_row,
        f"D{section_row + 3}",
        "Visão Geral dos Planos de Ação",
        status_labels,
        status_values,
        [STATUS_COLORS[label] for label in status_labels],
        width=12.5,
        height=7.2,
    )

    origin = indicators.get("visao_geral_origem", {})
    origin_labels = list(origin.keys())
    origin_values = [int(origin[label].get("quantidade", 0) or 0) for label in origin_labels]
    origin_data = [["Origem", "Quant."]]
    origin_data.extend([[label, value] for label, value in zip(origin_labels, origin_values)])
    origin_data.append(["TOTAL", sum(origin_values)])
    _write_excel_table(
        sheet,
        section_row + 3,
        11,
        origin_data,
        category_colors=ORIGIN_COLORS,
        light_header=True,
    )
    helper_row = _add_pie_chart(
        sheet,
        helper,
        helper_row,
        f"N{section_row + 3}",
        "Visão Geral de Origem",
        origin_labels,
        origin_values,
        [ORIGIN_COLORS[label] for label in origin_labels],
        width=12.5,
        height=7.2,
    )

    directorate_rows = indicators.get("planos_por_diretoria", [])
    section_row += 19
    print_break_rows.append(section_row)
    _section_title(sheet, section_row, "Planos de Ação por Diretoria")
    section_height = max(19, len(directorate_rows) + 5)
    _style_panel_range(sheet, section_row + 1, section_row + section_height, 1, 19)
    directorate_data = _group_export_table(directorate_rows, "Diretoria")
    _write_excel_table(
        sheet,
        section_row + 2,
        1,
        directorate_data,
        status_header=True,
        light_header=True,
    )
    helper_row = _add_group_chart(
        sheet,
        helper,
        helper_row,
        f"J{section_row + 2}",
        "Gráfico de Planos por Diretoria",
        directorate_rows,
        "Diretoria",
    )

    section_row += section_height + 3
    print_break_rows.append(section_row)
    unit_rows = indicators.get("planos_por_unidade", [])
    _section_title(sheet, section_row, "Planos de Ação por Unidade/Agência")
    section_height = max(21, len(unit_rows) + 5)
    _style_panel_range(sheet, section_row + 1, section_row + section_height, 1, 19)
    unit_data = _group_export_table(unit_rows, "Unidade/Agência")
    _write_excel_table(
        sheet,
        section_row + 2,
        1,
        unit_data,
        status_header=True,
        light_header=True,
    )
    helper_row = _add_group_chart(
        sheet,
        helper,
        helper_row,
        f"J{section_row + 2}",
        "Gráfico de Planos por Unidade/Agência",
        unit_rows,
        "Unidade/Agência",
        height=11.2,
    )

    section_row += section_height + 3
    print_break_rows.append(section_row)
    _section_title(sheet, section_row, "Repactuações e Planos Ativos")

    _style_panel_range(sheet, section_row + 1, section_row + 8, 1, 19)
    _excel_kpi_card(
        sheet,
        section_row + 2,
        section_row + 6,
        3,
        8,
        "PLANOS REPACTUADOS",
        int(indicators.get("planos_repactuados", 0) or 0),
        "#FFF200",
    )
    rate = float(indicators.get("taxa_repactuacao", 0) or 0)
    _excel_kpi_card(
        sheet,
        section_row + 2,
        section_row + 6,
        12,
        17,
        "TAXA DE REPACTUAÇÃO",
        f"{rate:.2f}%".replace(".", ","),
        "#FFF200",
    )

    content_row = section_row + 9
    _style_panel_range(sheet, content_row, content_row + 19, 1, 19)
    distribution = indicators.get("distribuicao_repactuacoes", {})
    distribution_rows = [
        ("1", int(distribution.get("1", 0) or 0)),
        ("2", int(distribution.get("2", 0) or 0)),
        ("3", int(distribution.get("3", 0) or 0)),
        ("Acima 3", int(distribution.get("Acima de 3", 0) or 0)),
    ]
    distribution_total = sum(value for _, value in distribution_rows)
    distribution_data = [["Repactuações", "Quant. Planos", "%"]]
    for label, value in distribution_rows:
        distribution_data.append(
            [label, value, (value / distribution_total) if distribution_total else 0]
        )
    distribution_data.append(
        ["TOTAL", distribution_total, 1 if distribution_total else 0]
    )
    _write_excel_table(
        sheet,
        content_row + 2,
        1,
        distribution_data,
        repactuation_rows=True,
    )
    for row in range(content_row + 3, content_row + 8):
        sheet.cell(row, 3).number_format = "0.0%"

    active_labels = ["Planos repactuados", "Demais planos ativos"]
    active_values = [
        int(indicators.get("planos_repactuados", 0) or 0),
        max(
            0,
            int(indicators.get("planos_ativos", 0) or 0)
            - int(indicators.get("planos_repactuados", 0) or 0),
        ),
    ]
    helper_row = _add_pie_chart(
        sheet,
        helper,
        helper_row,
        f"H{content_row + 1}",
        "Gráfico de Repactuações e Planos Ativos",
        active_labels,
        active_values,
        [ACTIVE_PLAN_COLORS[label] for label in active_labels],
        doughnut=True,
        width=17.8,
        height=9.3,
    )

    section_row = content_row + 22
    print_break_rows.append(section_row)
    period_rows = indicators.get(
        "periodos_conclusao_planos",
        indicators.get("repactuacoes_por_periodo", []),
    )
    _section_title(sheet, section_row, "Período Previsto para Conclusão de Planos")
    section_height = max(21, len(period_rows) + 5)
    _style_panel_range(sheet, section_row + 1, section_row + section_height, 1, 19)
    period_data = [["Planos Repactuados e Vencidos", "Quant. Planos"]]
    period_data.extend(
        [
            row.get("periodo", ""),
            int(row.get("quantidade", 0) or 0),
        ]
        for row in period_rows
    )
    _write_excel_table(
        sheet,
        section_row + 2,
        1,
        period_data,
        period_rows=True,
    )
    _add_period_chart(
        sheet,
        helper,
        helper_row,
        f"H{section_row + 2}",
        period_rows,
        height=11.2,
    )

    sheet.auto_filter.ref = f"A{section_row + 2}:B{section_row + len(period_data) + 1}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.outlinePr.summaryBelow = True
    sheet.print_area = f"A1:S{section_row + section_height}"
    for break_row in print_break_rows:
        sheet.row_breaks.append(Break(id=break_row - 1))
    sheet.freeze_panes = "A4"

    workbook.save(filename)
    return filename


def _semaforo_rules(indicators):
    return [
        {
            "condicao": str(rule.get("condicao", "")),
            "categoria": str(rule.get("categoria", "")),
            "quantidade": int(rule.get("quantidade", 0) or 0),
            "percentual": float(rule.get("percentual", 0) or 0),
            "planos": [
                str(plan)
                for plan in rule.get("relatorios", rule.get("planos", []))
            ],
        }
        for rule in indicators.get("regras", [])
    ]


def _semaforo_bar_chart_image(indicators):
    distribution = indicators.get("distribuicao_risco", {})
    labels = ["Alto", "Significativo", "Moderado", "Baixo"]
    values = [
        int((distribution.get(label, {}) or {}).get("quantidade", 0) or 0)
        for label in labels
    ]
    palette = [SEMAFORO_RISK_COLORS[label] for label in labels]
    fig, axis = plt.subplots(figsize=(9.4, 5.6))
    positions = list(range(len(labels)))
    bars = axis.barh(
        positions,
        values,
        color=palette,
        edgecolor="white",
        linewidth=1.0,
    )
    axis.set_yticks(positions, labels)
    axis.tick_params(axis="y", labelsize=9)
    axis.tick_params(axis="x", labelsize=8)
    axis.set_xlabel("Quantidade de planos", fontsize=9, fontweight="bold")
    axis.grid(axis="x", alpha=0.18)
    axis.set_axisbelow(True)
    axis.invert_yaxis()
    maximum = max(values, default=0)
    axis.set_xlim(0, max(1, maximum + 1))

    for bar, value in zip(bars, values):
        axis.text(
            max(0.04, bar.get_width() + 0.05),
            bar.get_y() + (bar.get_height() / 2),
            str(value),
            va="center",
            ha="left",
            fontsize=8,
            fontweight="bold",
            color="#111111",
        )

    fig.tight_layout()
    buffer = BytesIO()
    fig.savefig(buffer, format="png", dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    buffer.seek(0)
    return buffer


def _semaforo_pie_chart_image(indicators):
    distribution = indicators.get("distribuicao_risco", {})
    labels = ["Alto", "Significativo", "Moderado", "Baixo"]
    values = [
        int((distribution.get(label, {}) or {}).get("quantidade", 0) or 0)
        for label in labels
    ]
    palette = [SEMAFORO_RISK_COLORS[label] for label in labels]
    total = sum(values)
    fig, axis = plt.subplots(figsize=(9.4, 5.6))
    fig.subplots_adjust(bottom=0.19)

    def label_percentage(percent):
        if percent <= 0 or not total:
            return ""
        formatted = f"{percent:.1f}".replace(".", ",")
        return f"{formatted}%"

    if total:
        axis.pie(
            values,
            colors=palette,
            startangle=90,
            autopct=label_percentage,
            pctdistance=0.74,
            textprops={"fontsize": 8, "fontweight": "bold", "color": "#111111"},
            wedgeprops={"edgecolor": "white", "linewidth": 2},
        )
    else:
        axis.pie(
            [1],
            colors=["#E7EDF5"],
            startangle=90,
            wedgeprops={"edgecolor": "white", "linewidth": 2},
        )

    axis.axis("equal")
    handles = [
        Patch(facecolor=color, edgecolor="white")
        for color in palette
    ]
    axis.legend(
        handles,
        labels,
        loc="lower center",
        bbox_to_anchor=(0.5, -0.14),
        ncol=2,
        fontsize=8,
        frameon=True,
        facecolor="white",
        edgecolor="#D8E2EF",
    )
    buffer = BytesIO()
    fig.savefig(buffer, format="png", dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    buffer.seek(0)
    return buffer


def _semaforo_condition_chart_image(indicators):
    rules = _semaforo_rules(indicators)
    fig, axis = plt.subplots(figsize=(9.4, 5.6))
    axis.set_xlim(0, 1)
    axis.set_ylim(0, max(1, len(rules)))
    axis.axis("off")
    for index, rule in enumerate(rules):
        value = int(rule.get("quantidade", 0) or 0)
        color = SEMAFORO_CATEGORY_COLORS.get(rule.get("categoria"), BLUE_HEX)
        y = len(rules) - index - 0.5
        axis.scatter([0.08], [y], s=760, c=[color], edgecolors="#CBD6E4", linewidths=1.2)
        axis.text(
            0.08,
            y,
            f"({value})",
            ha="center",
            va="center",
            fontsize=7.5,
            fontweight="bold",
            color="#111111",
        )
        axis.text(
            0.15,
            y,
            fill(str(rule.get("condicao", "")), width=68),
            ha="left",
            va="center",
            fontsize=8.5,
            fontweight="bold",
            color="#111111",
        )
    fig.subplots_adjust(left=0.02, right=0.98, top=0.98, bottom=0.02)
    buffer = BytesIO()
    fig.savefig(buffer, format="png", dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    buffer.seek(0)
    return buffer


def _semaforo_kpi_cards_pdf(indicators):
    cards = []
    for category in ["Atenção", "Monitoramento", "Conhecimento"]:
        values = indicators.get("condicoes", {}).get(category, {})
        quantity = int(values.get("quantidade", 0) or 0)
        percentage = float(values.get("percentual", 0) or 0)
        percentage_text = f"{percentage:.2f}% do Total".replace(".", ",")
        cards.append(
            [
                _paragraph(category.upper(), size=12, bold=True, align=1),
                Spacer(1, 4),
                _paragraph(quantity, size=24, bold=True, align=1),
                Spacer(1, 3),
                _paragraph(percentage_text, size=9, bold=True, align=1),
            ]
        )

    table = Table(
        [cards],
        colWidths=[128 * mm] * 3,
        rowHeights=[30 * mm],
        hAlign="CENTER",
    )
    commands = [
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#B8C7D9")),
        ("INNERGRID", (0, 0), (-1, -1), 6, colors.white),
        ("ROUNDEDCORNERS", [9]),
        ("LEFTPADDING", (0, 0), (-1, -1), 12),
        ("RIGHTPADDING", (0, 0), (-1, -1), 12),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]
    for column, category in enumerate(["Atenção", "Monitoramento", "Conhecimento"]):
        commands.append(
            (
                "BACKGROUND",
                (column, 0),
                (column, 0),
                colors.HexColor(SEMAFORO_CATEGORY_COLORS[category]),
            )
        )
    table.setStyle(TableStyle(commands))
    return table


def _semaforo_rules_pdf_table(indicators):
    rules = _semaforo_rules(indicators)
    data = [[
        "Condição",
        "Categoria",
        "Quantidade",
        "% do Total",
        "Relatórios de Auditoria enquadrados",
    ]]
    data.extend(
        [
            _paragraph(rule["condicao"], size=6.5),
            rule["categoria"],
            rule["quantidade"],
            f'{rule["percentual"]:.2f}%'.replace(".", ","),
            _paragraph(
                "; ".join(
                    f"{index}. {plan}"
                    for index, plan in enumerate(rule["planos"], start=1)
                ) or "Nenhum plano enquadrado nesta condição.",
                size=6,
            ),
        ]
        for rule in rules
    )
    data.append(
        [
            _paragraph("TOTAL", size=7, bold=True),
            "",
            sum(rule["quantidade"] for rule in rules),
            "100,00%" if sum(rule["quantidade"] for rule in rules) else "0,00%",
            "",
        ]
    )
    table = Table(
        data,
        colWidths=[84 * mm, 43 * mm, 28 * mm, 26 * mm, 201 * mm],
        repeatRows=1,
    )
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), BLUE),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("BACKGROUND", (0, -1), (-1, -1), BLUE),
        ("TEXTCOLOR", (0, -1), (-1, -1), colors.white),
        ("ALIGN", (1, 0), (3, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTSIZE", (0, 0), (-1, -1), 6.5),
        ("GRID", (0, 0), (-1, -1), 0.35, THIN_GREY),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    for row, rule in enumerate(rules, start=1):
        background = SEMAFORO_CATEGORY_COLORS.get(rule["categoria"], LIGHT_BLUE_HEX)
        commands.extend(
            [
                ("BACKGROUND", (1, row), (1, row), colors.HexColor(background)),
                ("FONTNAME", (1, row), (1, row), "Helvetica-Bold"),
                ("TEXTCOLOR", (1, row), (1, row), colors.HexColor("#111111")),
            ]
        )
    table.setStyle(TableStyle(commands))
    return table


def _reincidence_data(indicators):
    return indicators.get("reincidencia", {})


def _reincidence_kpi_cards_pdf(indicators):
    recurrence = _reincidence_data(indicators)
    rate = float(recurrence.get("taxa_reincidencia", 0) or 0)
    cards = [
        (
            "PLANOS EM MONITORAMENTO",
            int(recurrence.get("planos_monitoramento", 0) or 0),
            REINCIDENCE_KPI_COLORS["Planos em Monitoramento"],
        ),
        (
            "PLANOS REINCIDENTES",
            int(recurrence.get("planos_reincidentes", 0) or 0),
            REINCIDENCE_KPI_COLORS["Planos Reincidentes"],
        ),
        (
            "TAXA DE REINCIDÊNCIA",
            f"{rate:.2f}%".replace(".", ","),
            REINCIDENCE_KPI_COLORS["Taxa de Reincidência"],
        ),
    ]
    contents = [
        [
            _paragraph(label, size=11, bold=True, align=1),
            Spacer(1, 5),
            _paragraph(value, size=25, bold=True, align=1),
        ]
        for label, value, _color in cards
    ]
    table = Table(
        [contents],
        colWidths=[128 * mm] * 3,
        rowHeights=[30 * mm],
        hAlign="CENTER",
    )
    commands = [
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#B8C7D9")),
        ("INNERGRID", (0, 0), (-1, -1), 6, colors.white),
        ("ROUNDEDCORNERS", [9]),
        ("LEFTPADDING", (0, 0), (-1, -1), 12),
        ("RIGHTPADDING", (0, 0), (-1, -1), 12),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]
    for column, (_label, _value, color) in enumerate(cards):
        commands.append(
            (
                "BACKGROUND",
                (column, 0),
                (column, 0),
                colors.HexColor(color),
            )
        )
    table.setStyle(TableStyle(commands))
    return table


def _reincidence_criticality_chart_image(indicators):
    recurrence = _reincidence_data(indicators)
    criticality = recurrence.get("criticidade_reincidente", {})
    labels = ["Alto", "Significativo", "Moderado", "Baixo"]
    values = [int(criticality.get(label, 0) or 0) for label in labels]
    palette = [REINCIDENCE_RISK_COLORS[label] for label in labels]
    fig, axis = plt.subplots(figsize=(9.4, 5.6))
    positions = list(range(len(labels)))
    bars = axis.barh(
        positions,
        values,
        color=palette,
        edgecolor="white",
        linewidth=1.0,
    )
    axis.set_yticks(positions, labels)
    axis.tick_params(axis="y", labelsize=9)
    axis.tick_params(axis="x", labelsize=8)
    axis.set_xlabel("Quantidade de planos", fontsize=9, fontweight="bold")
    axis.grid(axis="x", alpha=0.18)
    axis.set_axisbelow(True)
    axis.invert_yaxis()
    maximum = max(values, default=0)
    axis.set_xlim(0, max(1, maximum + 1))

    for bar, value in zip(bars, values):
        axis.text(
            max(0.04, bar.get_width() + 0.05),
            bar.get_y() + (bar.get_height() / 2),
            str(value),
            va="center",
            ha="left",
            fontsize=9,
            fontweight="bold",
            color="#111111",
        )

    fig.tight_layout()
    buffer = BytesIO()
    fig.savefig(buffer, format="png", dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    buffer.seek(0)
    return buffer


def _reincidence_reports_chart_image(indicators):
    recurrence = _reincidence_data(indicators)
    reports = recurrence.get("reincidentes_por_relatorio", [])
    fig, axis = plt.subplots(figsize=(9.4, 5.6))
    axis.set_xlim(0, 1)
    axis.set_ylim(0, max(1, len(reports)))
    axis.axis("off")
    for index, report in enumerate(reports):
        value = int(report.get("quantidade", 0) or 0)
        color = REINCIDENCE_REPORT_COLORS[index % len(REINCIDENCE_REPORT_COLORS)]
        y = len(reports) - index - 0.5
        axis.scatter([0.08], [y], s=760, c=[color], edgecolors="#CBD6E4", linewidths=1.2)
        axis.text(
            0.08,
            y,
            f"({value})",
            ha="center",
            va="center",
            fontsize=7.5,
            fontweight="bold",
            color="#111111",
        )
        axis.text(
            0.15,
            y,
            fill(str(report.get("relatorio", "")), width=68),
            ha="left",
            va="center",
            fontsize=max(6.4, 8.5 - max(0, len(reports) - 5) * 0.3),
            fontweight="bold",
            color="#111111",
        )
    if not reports:
        axis.text(
            0.5,
            0.5,
            "Sem dados de reincidência.",
            transform=axis.transAxes,
            ha="center",
            va="center",
            fontsize=10,
            fontweight="bold",
            color="#52647A",
        )
    fig.subplots_adjust(left=0.02, right=0.98, top=0.98, bottom=0.02)
    buffer = BytesIO()
    fig.savefig(buffer, format="png", dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    buffer.seek(0)
    return buffer


def _reincidence_criticality_pdf_table(indicators):
    recurrence = _reincidence_data(indicators)
    criticality = recurrence.get("criticidade_reincidente", {})
    labels = ["Alto", "Significativo", "Moderado", "Baixo"]
    data = [["Criticidade", "Quantidade"]]
    data.extend(
        [label, int(criticality.get(label, 0) or 0)]
        for label in labels
    )
    data.append(["TOTAL", sum(row[1] for row in data[1:])])
    table = Table(data, colWidths=[76 * mm, 35 * mm], rowHeights=8 * mm)
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), BLUE),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, -1), (-1, -1), BLUE),
        ("TEXTCOLOR", (0, -1), (-1, -1), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.35, THIN_GREY),
    ]
    for row, label in enumerate(labels, start=1):
        commands.extend(
            [
                ("BACKGROUND", (0, row), (0, row), colors.HexColor(REINCIDENCE_RISK_COLORS[label])),
                ("FONTNAME", (0, row), (0, row), "Helvetica-Bold"),
                ("TEXTCOLOR", (0, row), (0, row), colors.HexColor("#111111")),
            ]
        )
    table.setStyle(TableStyle(commands))
    return table


def _reincidence_reports_pdf_table(indicators):
    recurrence = _reincidence_data(indicators)
    reports = recurrence.get("reincidentes_por_relatorio", [])
    data = [["Relatório de Auditoria", "Quantidade", "%"]]
    data.extend(
        [
            _paragraph(item.get("relatorio", ""), size=6.5),
            int(item.get("quantidade", 0) or 0),
            f'{float(item.get("percentual", 0) or 0):.2f}%'.replace(".", ","),
        ]
        for item in reports
    )
    total = sum(int(item.get("quantidade", 0) or 0) for item in reports)
    data.append([_paragraph("TOTAL", size=6.5, bold=True), total, "100,00%" if total else "0,00%"])
    table = Table(data, colWidths=[132 * mm, 30 * mm, 25 * mm], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), BLUE),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("BACKGROUND", (0, -1), (-1, -1), BLUE),
                ("TEXTCOLOR", (0, -1), (-1, -1), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, LIGHT_BLUE]),
                ("ALIGN", (1, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.35, THIN_GREY),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def _active_plans_pdf_elements(active_plans, title_style, subtitle_style):
    if active_plans is None or active_plans.empty:
        return []
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle(
        "ActivePlanCell",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=5.5,
        leading=6.5,
    )
    columns = active_plans.columns.tolist()
    chunk_size = 6
    elements = []
    for start in range(0, len(columns), chunk_size):
        chunk = columns[start:start + chunk_size]
        elements.extend(
            [
                PageBreak(),
                _title_bar(f"Tabela Analítica - {len(active_plans)} Planos Ativos"),
                Spacer(1, 3 * mm),
                Paragraph(
                    escape(
                        f"Somente planos ativos | Total: {len(active_plans)} | "
                        f"Colunas {start + 1} a {start + len(chunk)} de {len(columns)}"
                    ),
                    subtitle_style,
                ),
                Spacer(1, 5 * mm),
            ]
        )
        data = [
            [Paragraph("Registro", cell_style)]
            + [Paragraph(escape(str(column)), cell_style) for column in chunk]
        ]
        for record, (_index, row) in enumerate(active_plans[chunk].iterrows(), start=1):
            data.append(
                [Paragraph(str(record), cell_style)]
                + [
                    Paragraph(
                        escape(_safe_text(row[column])).replace("\n", "<br/>"),
                        cell_style,
                    )
                    for column in chunk
                ]
            )
        available = landscape(A3)[0] - (20 * mm)
        elements.append(
            _styled_table(
                data,
                [18 * mm] + [(available - (18 * mm)) / len(chunk)] * len(chunk),
                5.5,
            )
        )
    return elements


def export_semaforo_pdf(indicators, active_plans=None, filters=None):
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    filename = EXPORT_DIR / f"semaforo_planos_acao_{_timestamp()}.pdf"
    document = SimpleDocTemplate(
        str(filename),
        pagesize=landscape(A3),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
        title="Semáforo de Planos de Ação",
        author="Dashboard de Monitoramento",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "SemaforoTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=21,
        textColor=colors.HexColor("#0E2146"),
        alignment=1,
        spaceAfter=3,
    )
    subtitle_style = ParagraphStyle(
        "SemaforoSubtitle",
        parent=styles["BodyText"],
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#52647A"),
        alignment=1,
    )
    reference = indicators.get("data_referencia", "")
    filters = filters or {}
    total = int(indicators.get("total_planos_ativos", 0) or 0)
    classified = int(indicators.get("total_classificados", 0) or 0)
    elements = [
        Paragraph("Semáforo de Planos de Ação", title_style),
        Paragraph(
            escape(
                f"Data de referência: {reference} | Planos ativos: {total} | "
                f"Total classificado: {classified}"
            ),
            subtitle_style,
        ),
        Paragraph(
            escape(f"Filtros: {_filter_text(filters)}"),
            subtitle_style,
        ),
        Spacer(1, 5 * mm),
        _semaforo_kpi_cards_pdf(indicators),
        Spacer(1, 5 * mm),
    ]
    charts = Table(
        [
            [
                _panel_pdf(
                    "Planos Ativos por Risco",
                    Image(_semaforo_bar_chart_image(indicators), width=118 * mm, height=92 * mm),
                    126 * mm,
                ),
                _panel_pdf(
                    "Distribuição por Risco",
                    Image(_semaforo_pie_chart_image(indicators), width=118 * mm, height=92 * mm),
                    126 * mm,
                ),
                _panel_pdf(
                    "Condições do Semáforo",
                    Image(
                        _semaforo_condition_chart_image(indicators),
                        width=118 * mm,
                        height=92 * mm,
                    ),
                    126 * mm,
                ),
            ]
        ],
        colWidths=[130 * mm, 130 * mm, 130 * mm],
    )
    charts.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 2),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    elements.extend(
        [
            charts,
            Spacer(1, 4 * mm),
            _semaforo_rules_pdf_table(indicators),
        ]
    )
    recurrence = _reincidence_data(indicators)
    active_count = int(recurrence.get("planos_monitoramento", 0) or 0)
    recurrent_plans = int(recurrence.get("planos_reincidentes", 0) or 0)
    recurrence_rate = float(recurrence.get("taxa_reincidencia", 0) or 0)
    elements.extend(
        [
            PageBreak(),
            Paragraph("Indicadores de Reincidência", title_style),
            Paragraph(
                escape(
                    (
                        f"Planos ativos: {active_count} | "
                        f"Planos reincidentes: {recurrent_plans} | "
                        f"Taxa de reincidência: {recurrence_rate:.2f}%"
                    ).replace(".", ",")
                ),
                subtitle_style,
            ),
            Spacer(1, 5 * mm),
            _reincidence_kpi_cards_pdf(indicators),
            Spacer(1, 5 * mm),
        ]
    )
    recurrence_charts = Table(
        [
            [
                _panel_pdf(
                    "Criticidade de Recomendação Reincidente",
                    Image(
                        _reincidence_criticality_chart_image(indicators),
                        width=184 * mm,
                        height=101 * mm,
                    ),
                    194 * mm,
                ),
                _panel_pdf(
                    "Recomendação Reincidente de Auditoria",
                    Image(
                        _reincidence_reports_chart_image(indicators),
                        width=184 * mm,
                        height=101 * mm,
                    ),
                    194 * mm,
                ),
            ]
        ],
        colWidths=[196 * mm, 196 * mm],
    )
    recurrence_charts.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 2),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    recurrence_tables = Table(
        [
            [
                _panel_pdf(
                    "Criticidade dos Planos Reincidentes",
                    _reincidence_criticality_pdf_table(indicators),
                    194 * mm,
                ),
                _panel_pdf(
                    "Planos Reincidentes por Relatório de Auditoria",
                    _reincidence_reports_pdf_table(indicators),
                    194 * mm,
                ),
            ]
        ],
        colWidths=[196 * mm, 196 * mm],
    )
    recurrence_tables.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 2),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    elements.extend(
        [
            recurrence_charts,
            Spacer(1, 4 * mm),
            recurrence_tables,
        ]
    )
    elements.extend(_active_plans_pdf_elements(active_plans, title_style, subtitle_style))
    document.build(
        elements,
        onFirstPage=_pdf_page_footer,
        onLaterPages=_pdf_page_footer,
    )
    return filename


def _add_semaforo_rules_excel_sheet(workbook, indicators):
    sheet = workbook.create_sheet("Condições")
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A5"
    sheet.merge_cells("A1:E1")
    sheet["A1"] = "Condições do Semáforo e Planos de Ação"
    sheet["A1"].font = Font(size=17, bold=True, color="0E2146")
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.merge_cells("A2:E2")
    sheet["A2"] = "Cada plano está listado individualmente na condição em que foi enquadrado."
    sheet["A2"].font = Font(size=9, italic=True, color="52647A")
    sheet["A2"].alignment = Alignment(horizontal="center", vertical="center")
    headers = [
        "Condição",
        "Categoria",
        "Quantidade",
        "% do Total",
        "Relatórios de Auditoria enquadrados",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(4, column, header)
        cell.fill = _excel_fill(BLUE_HEX)
        cell.font = Font(size=9, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            left=Side(style="thin", color="9BAAC2"),
            right=Side(style="thin", color="9BAAC2"),
            top=Side(style="thin", color="9BAAC2"),
            bottom=Side(style="thin", color="9BAAC2"),
        )
    rules = _semaforo_rules(indicators)
    for row, rule in enumerate(rules, start=5):
        plans = "\n".join(
            f"{index}. {plan}"
            for index, plan in enumerate(rule["planos"], start=1)
        ) or "Nenhum plano enquadrado nesta condição."
        values = [
            rule["condicao"],
            rule["categoria"],
            rule["quantidade"],
            rule["percentual"] / 100,
            plans,
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row, column, value)
            cell.alignment = Alignment(
                horizontal="center" if column in {2, 3, 4} else "left",
                vertical="top",
                wrap_text=True,
            )
            cell.fill = _excel_fill("FFFFFF" if row % 2 else LIGHT_BLUE_HEX)
            cell.border = Border(
                left=Side(style="thin", color="9BAAC2"),
                right=Side(style="thin", color="9BAAC2"),
                top=Side(style="thin", color="9BAAC2"),
                bottom=Side(style="thin", color="9BAAC2"),
            )
        sheet.cell(row, 2).fill = _excel_fill(
            SEMAFORO_CATEGORY_COLORS.get(rule["categoria"], LIGHT_BLUE_HEX)
        )
        sheet.cell(row, 2).font = Font(bold=True, color="111111")
        sheet.cell(row, 4).number_format = "0.00%"
        estimated_lines = sum(
            max(1, (len(plan) + 89) // 90)
            for plan in rule["planos"]
        ) or 1
        sheet.row_dimensions[row].height = max(30, (estimated_lines * 15) + 8)
    total_row = 5 + len(rules)
    totals = [
        "TOTAL",
        "",
        sum(rule["quantidade"] for rule in rules),
        1 if sum(rule["quantidade"] for rule in rules) else 0,
        "",
    ]
    for column, value in enumerate(totals, start=1):
        cell = sheet.cell(total_row, column, value)
        cell.fill = _excel_fill(BLUE_HEX)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.cell(total_row, 4).number_format = "0.00%"
    sheet.column_dimensions["A"].width = 58
    sheet.column_dimensions["B"].width = 22
    sheet.column_dimensions["C"].width = 14
    sheet.column_dimensions["D"].width = 15
    sheet.column_dimensions["E"].width = 90
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[4].height = 30
    sheet.auto_filter.ref = f"A4:E{total_row - 1}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:4"
    sheet.print_area = f"A1:E{total_row}"


def _add_active_plans_excel_sheet(workbook, active_plans):
    if active_plans is None or active_plans.empty:
        return
    sheet = workbook.create_sheet("Planos Ativos")
    sheet.sheet_view.showGridLines = False
    last_column = get_column_letter(len(active_plans.columns))
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(active_plans.columns))
    sheet["A1"] = f"Tabela Analítica - {len(active_plans)} Planos Ativos"
    sheet["A1"].fill = _excel_fill(BLUE_HEX)
    sheet["A1"].font = Font(size=17, bold=True, color="FFFFFF")
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(active_plans.columns))
    sheet["A2"] = f"Somente planos ativos | Total: {len(active_plans)}"
    sheet["A2"].font = Font(size=9, italic=True, color="52647A")
    sheet["A2"].alignment = Alignment(horizontal="center", vertical="center")
    for column, name in enumerate(active_plans.columns, start=1):
        cell = sheet.cell(4, column, str(name))
        cell.fill = _excel_fill(BLUE_HEX)
        cell.font = Font(size=9, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_number, values in enumerate(active_plans.itertuples(index=False, name=None), start=5):
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column, _safe_text(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.fill = _excel_fill("FFFFFF" if row_number % 2 else LIGHT_BLUE_HEX)
            cell.border = Border(
                left=Side(style="thin", color="D8E2EF"),
                right=Side(style="thin", color="D8E2EF"),
                top=Side(style="thin", color="D8E2EF"),
                bottom=Side(style="thin", color="D8E2EF"),
            )
        sheet.row_dimensions[row_number].height = 36
    for column, name in enumerate(active_plans.columns, start=1):
        values = [_safe_text(name)] + [
            _safe_text(value) for value in active_plans.iloc[:, column - 1]
        ]
        sheet.column_dimensions[get_column_letter(column)].width = max(
            12,
            min(45, max(len(value) for value in values) + 2),
        )
    last_row = 4 + len(active_plans)
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:{last_column}{last_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:4"
    sheet.print_area = f"A1:{last_column}{last_row}"


def export_semaforo_excel(indicators, active_plans=None, filters=None):
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    filename = EXPORT_DIR / f"semaforo_planos_acao_{_timestamp()}.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Semáforo"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"

    for column in range(1, 20):
        sheet.column_dimensions[get_column_letter(column)].width = (
            17 if column in {1, 6, 11, 16} else 11.5
        )
    for row in range(1, 86):
        sheet.row_dimensions[row].height = 20

    _paint_range(sheet, 1, 85, 1, 19, _excel_fill(PAGE_BACKGROUND_HEX))
    _merge_title(sheet, 1, "Semáforo de Planos de Ação")
    sheet.merge_cells("A3:S3")
    filters = filters or {}
    reference = indicators.get("data_referencia", "")
    total = int(indicators.get("total_planos_ativos", 0) or 0)
    classified = int(indicators.get("total_classificados", 0) or 0)
    sheet["A3"] = (
        f"Data de referência: {reference} | Planos ativos: {total} | "
        f"Total classificado: {classified} | Filtros: {_filter_text(filters)}"
    )
    sheet["A3"].font = Font(size=9, italic=True, color="52647A")
    sheet["A3"].alignment = Alignment(horizontal="center", vertical="center")

    for first_col, last_col, category in [
        (1, 5, "Atenção"),
        (8, 12, "Monitoramento"),
        (15, 19, "Conhecimento"),
    ]:
        values = indicators.get("condicoes", {}).get(category, {})
        quantity = int(values.get("quantidade", 0) or 0)
        percentage = float(values.get("percentual", 0) or 0)
        value = f'{quantity}\n{percentage:.2f}% do Total'.replace(".", ",")
        _excel_kpi_card(
            sheet,
            5,
            8,
            first_col,
            last_col,
            category.upper(),
            value,
            SEMAFORO_CATEGORY_COLORS[category],
        )

    _section_title(sheet, 10, "Indicadores por Risco e Condição")
    _style_panel_range(sheet, 11, 31, 1, 19)
    chart_titles = [
        (1, 6, "Planos Ativos por Risco"),
        (7, 12, "Distribuição por Risco"),
        (13, 19, "Condições do Semáforo"),
    ]
    for first_col, last_col, title in chart_titles:
        sheet.merge_cells(
            start_row=11,
            start_column=first_col,
            end_row=11,
            end_column=last_col,
        )
        cell = sheet.cell(11, first_col, title)
        cell.font = Font(size=11, bold=True, color="111111")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    chart_images = [
        (XLImage(_semaforo_bar_chart_image(indicators)), "A12"),
        (XLImage(_semaforo_pie_chart_image(indicators)), "G12"),
        (XLImage(_semaforo_condition_chart_image(indicators)), "M12"),
    ]
    for chart, anchor in chart_images:
        chart.width = int(12 * 37.8)
        chart.height = int(9.8 * 37.8)
        sheet.add_image(chart, anchor)

    recurrence = _reincidence_data(indicators)
    _section_title(sheet, 33, "Indicadores de Reincidência")
    _style_panel_range(sheet, 34, 39, 1, 19)
    rate = float(recurrence.get("taxa_reincidencia", 0) or 0)
    recurrence_cards = [
        (
            1,
            5,
            "PLANOS EM MONITORAMENTO",
            int(recurrence.get("planos_monitoramento", 0) or 0),
            REINCIDENCE_KPI_COLORS["Planos em Monitoramento"],
        ),
        (
            8,
            12,
            "PLANOS REINCIDENTES",
            int(recurrence.get("planos_reincidentes", 0) or 0),
            REINCIDENCE_KPI_COLORS["Planos Reincidentes"],
        ),
        (
            15,
            19,
            "TAXA DE REINCIDÊNCIA",
            f"{rate:.2f}%".replace(".", ","),
            REINCIDENCE_KPI_COLORS["Taxa de Reincidência"],
        ),
    ]
    for first_col, last_col, label, value, color in recurrence_cards:
        _excel_kpi_card(sheet, 35, 38, first_col, last_col, label, value, color)

    _section_title(sheet, 41, "Criticidade de Recomendação Reincidente")
    _style_panel_range(sheet, 42, 59, 1, 19)
    criticality = recurrence.get("criticidade_reincidente", {})
    risk_labels = ["Alto", "Significativo", "Moderado", "Baixo"]
    criticality_data = [["Criticidade", "Quantidade"]]
    criticality_data.extend(
        [label, int(criticality.get(label, 0) or 0)]
        for label in risk_labels
    )
    criticality_data.append(
        [
            "TOTAL",
            sum(int(criticality.get(label, 0) or 0) for label in risk_labels),
        ]
    )
    _write_excel_table(
        sheet,
        43,
        1,
        criticality_data,
        category_colors=REINCIDENCE_RISK_COLORS,
    )
    criticality_image = XLImage(_reincidence_criticality_chart_image(indicators))
    criticality_image.width = int(19.5 * 37.8)
    criticality_image.height = int(10.5 * 37.8)
    sheet.add_image(criticality_image, "F43")

    _section_title(sheet, 61, "Recomendação Reincidente de Auditoria")
    _style_panel_range(sheet, 62, 85, 1, 19)
    reports = recurrence.get("reincidentes_por_relatorio", [])
    report_data = [["Relatório de Auditoria", "Quantidade", "%"]]
    report_data.extend(
        [
            item.get("relatorio", ""),
            int(item.get("quantidade", 0) or 0),
            float(item.get("percentual", 0) or 0) / 100,
        ]
        for item in reports
    )
    report_total = sum(int(item.get("quantidade", 0) or 0) for item in reports)
    report_data.append(["TOTAL", report_total, 1 if report_total else 0])
    _write_excel_table(sheet, 63, 1, report_data)
    for row in range(64, 64 + len(reports) + 1):
        sheet.cell(row, 3).number_format = "0.00%"
    reports_image = XLImage(_reincidence_reports_chart_image(indicators))
    reports_image.width = int(20.5 * 37.8)
    reports_image.height = int(10.8 * 37.8)
    sheet.add_image(reports_image, "G63")

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.print_area = "A1:S85"
    sheet.row_breaks.append(Break(id=32))
    sheet.row_breaks.append(Break(id=60))

    _add_semaforo_rules_excel_sheet(workbook, indicators)
    _add_active_plans_excel_sheet(workbook, active_plans)
    workbook.save(filename)
    return filename


def _paint_process_chart_image(indicators):
    rows = indicators.get("processos_status", [])
    labels = [str(row.get("processo", "")) for row in rows]
    statuses = [str(row.get("status", "")) for row in rows]
    figure, axis = plt.subplots(figsize=(9.0, 5.8))

    positions = list(range(len(labels)))
    colors_list = [PAINT_STATUS_COLORS.get(status, BLUE_HEX) for status in statuses]
    axis.barh(
        positions,
        [0.66] * len(labels),
        color=colors_list,
        edgecolor="white",
        linewidth=1.2,
        height=0.68,
    )
    axis.set_yticks(positions)
    axis.set_yticklabels(
        labels,
        fontsize=9.2,
        fontweight="bold",
    )
    axis.set_xlim(0, 1)
    axis.set_xticks([])
    axis.grid(False)
    axis.set_title(
        "Processos e Status",
        fontsize=12,
        fontweight="bold",
        color="#111111",
        pad=10,
    )
    axis.invert_yaxis()
    for index, status in enumerate(statuses):
        axis.text(
            0.33,
            index,
            status,
            ha="center",
            va="center",
            fontsize=8,
            fontweight="bold",
            color=_text_color_for(PAINT_STATUS_COLORS.get(status, BLUE_HEX)),
        )
    axis.spines[["top", "right", "bottom", "left"]].set_visible(False)
    axis.legend(
        handles=[
            Patch(facecolor=color, edgecolor="white", label=label)
            for label, color in PAINT_STATUS_COLORS.items()
        ],
        loc="lower center",
        bbox_to_anchor=(0.5, -0.08),
        ncol=3,
        fontsize=8,
        frameon=True,
        facecolor="white",
        edgecolor="#D8E2EF",
    )
    figure.subplots_adjust(left=0.49, right=0.99, top=0.91, bottom=0.15)
    buffer = BytesIO()
    figure.savefig(buffer, format="png", dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(figure)
    buffer.seek(0)
    return buffer


def _paint_distribution_chart_image(indicators):
    distribution = indicators.get("distribuicao_status", {})
    labels = list(PAINT_STATUS_COLORS.keys())
    values = [int(distribution.get(label, {}).get("quantidade", 0) or 0) for label in labels]
    return _pie_chart_image(
        labels,
        values,
        "Gráfico por Status",
        PAINT_STATUS_COLORS,
        show_value_with_percent=True,
    )


def _paint_kpi_cards_pdf(indicators):
    cards = [
        ("TOTAL DE AUDITORIAS", int(indicators.get("total_auditorias", 0) or 0)),
        ("CONCLUÍDAS", int(indicators.get("concluidas", 0) or 0)),
        ("EM ANDAMENTO", int(indicators.get("em_andamento", 0) or 0)),
        ("NÃO INICIADAS", int(indicators.get("nao_iniciadas", 0) or 0)),
    ]
    contents = [
        [
            _paragraph(label, size=9, bold=True, align=1),
            Spacer(1, 4),
            _paragraph(value, size=24, bold=True, align=1),
        ]
        for label, value in cards
    ]
    table = Table(
        [contents],
        colWidths=[95.5 * mm] * 4,
        rowHeights=[30 * mm],
        hAlign="CENTER",
    )
    commands = [
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#B8C7D9")),
        ("INNERGRID", (0, 0), (-1, -1), 6, colors.white),
        ("ROUNDEDCORNERS", [9]),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]
    for column, color in enumerate(PAINT_KPI_COLORS):
        commands.append(("BACKGROUND", (column, 0), (column, 0), colors.HexColor(color)))
    table.setStyle(TableStyle(commands))
    return table


def _paint_source_text(source):
    details = [f"Base: {source.get('arquivo', 'execucao_paint.xlsx')}"]
    if source.get("periodo"):
        details.append(f"Período: {source['periodo']}")
    if source.get("posicao"):
        details.append(f"Posição: {source['posicao']}")
    return " | ".join(details)


def _paint_pdf_table(dataframe):
    styles = getSampleStyleSheet()
    header_style = ParagraphStyle(
        "PaintTableHeader",
        parent=styles["BodyText"],
        fontName="Helvetica-Bold",
        fontSize=7,
        leading=8,
        textColor=colors.white,
        alignment=1,
    )
    cell_style = ParagraphStyle(
        "PaintTableCell",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=6.4,
        leading=7.6,
        textColor=colors.HexColor("#111111"),
    )
    columns = dataframe.columns.tolist()
    data = [[Paragraph(escape(str(column)), header_style) for column in columns]]
    for _, row in dataframe.iterrows():
        data.append(
            [
                Paragraph(
                    escape(_safe_text(row[column])).replace("\n", "<br/>"),
                    cell_style,
                )
                for column in columns
            ]
        )

    available = landscape(A3)[0] - (20 * mm)
    width_map = {
        "Ano": 17 * mm,
        "N.": 12 * mm,
        "Auditoria": 82 * mm,
        "Processo": 132 * mm,
        "Realização": 82 * mm,
        "Status": 58 * mm,
    }
    widths = [width_map.get(column, available / max(1, len(columns))) for column in columns]
    table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), BLUE),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.35, THIN_GREY),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.white]),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    table.setStyle(TableStyle(commands))
    return table


def export_paint_pdf(indicators, dataframe, filters, source):
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    filename = EXPORT_DIR / f"execucao_paint_{_timestamp()}.pdf"
    document = _base_doc(filename)
    elements = []

    elements.extend(_pdf_header("Indicadores e visões", filters, title="Execução do PAINT"))
    elements.extend(
        [
            Paragraph(
                escape(_paint_source_text(source)),
                ParagraphStyle(
                    "PaintSource",
                    parent=getSampleStyleSheet()["BodyText"],
                    fontSize=8,
                    leading=10,
                    textColor=colors.HexColor("#52647A"),
                ),
            ),
            Spacer(1, 3 * mm),
            _paint_kpi_cards_pdf(indicators),
            Spacer(1, 6 * mm),
        ]
    )

    charts = Table(
        [
            [
                _panel_pdf(
                    "Processos e Status",
                    Image(_paint_process_chart_image(indicators), width=182 * mm, height=110 * mm),
                    194 * mm,
                ),
                _panel_pdf(
                    "Gráfico por Status",
                    Image(_paint_distribution_chart_image(indicators), width=182 * mm, height=110 * mm),
                    194 * mm,
                ),
            ]
        ],
        colWidths=[196 * mm, 196 * mm],
    )
    charts.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 2),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    elements.append(charts)

    elements.append(PageBreak())
    elements.extend(
        _pdf_header(
            "Tabela Analítica",
            filters,
            title="Execução do PAINT",
        )
    )
    elements.extend(
        [
            Paragraph(
                escape(_paint_source_text(source)),
                getSampleStyleSheet()["Normal"],
            ),
            Spacer(1, 3 * mm),
            _paint_pdf_table(dataframe),
        ]
    )

    document.build(
        elements,
        onFirstPage=_pdf_page_footer,
        onLaterPages=_pdf_page_footer,
    )
    return filename


def _write_paint_analytical_sheet(workbook, dataframe, filters):
    sheet = workbook.create_sheet("Tabela Analítica")
    sheet.sheet_view.showGridLines = False
    last_column = get_column_letter(max(1, len(dataframe.columns)))
    sheet.merge_cells(f"A1:{last_column}1")
    sheet["A1"] = "Tabela Analítica da Execução do PAINT"
    sheet["A1"].fill = _excel_fill(BLUE_HEX)
    sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet["A2"] = _filter_text(filters)
    sheet["A2"].font = Font(size=9, italic=True, color="52647A")
    sheet.row_dimensions[1].height = 28

    header_row = 4
    thin = Side(style="thin", color="9BAAC2")
    for column_index, column in enumerate(dataframe.columns, start=1):
        cell = sheet.cell(header_row, column_index, column)
        cell.fill = _excel_fill(BLUE_HEX)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row_index, (_, row) in enumerate(dataframe.iterrows(), start=header_row + 1):
        for column_index, column in enumerate(dataframe.columns, start=1):
            value = row[column]
            if pd.isna(value):
                value = ""
            elif hasattr(value, "item"):
                value = value.item()
            cell = sheet.cell(row_index, column_index, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.fill = _excel_fill("#FFFFFF")

    widths = {
        "Ano": 12,
        "N.": 8,
        "Auditoria": 38,
        "Processo": 52,
        "Realização": 32,
        "Status": 21,
    }
    for column_index, column in enumerate(dataframe.columns, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = widths.get(column, 24)
    last_row = header_row + len(dataframe)
    sheet.auto_filter.ref = f"A{header_row}:{last_column}{last_row}"
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:4"
    sheet.print_area = f"A1:{last_column}{last_row}"


def export_paint_table_pdf(dataframe, filters):
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    filename = EXPORT_DIR / f"tabela_execucao_paint_{_timestamp()}.pdf"
    document = _base_doc(filename)
    elements = []
    elements.extend(
        _pdf_header(
            "Tabela Analítica",
            filters,
            title="Execução do PAINT",
        )
    )
    elements.append(_paint_pdf_table(dataframe))
    document.build(
        elements,
        onFirstPage=_pdf_page_footer,
        onLaterPages=_pdf_page_footer,
    )
    return filename


def export_paint_table_excel(dataframe, filters):
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    filename = EXPORT_DIR / f"tabela_execucao_paint_{_timestamp()}.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_paint_analytical_sheet(workbook, dataframe, filters)
    workbook.save(filename)
    return filename


def export_paint_excel(indicators, dataframe, filters, source):
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    filename = EXPORT_DIR / f"execucao_paint_{_timestamp()}.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Execução do PAINT"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"

    for column in range(1, 20):
        sheet.column_dimensions[get_column_letter(column)].width = 11.5
    for row in range(1, 75):
        sheet.row_dimensions[row].height = 20
    _paint_range(sheet, 1, 75, 1, 19, _excel_fill(PAGE_BACKGROUND_HEX))

    _merge_title(sheet, 1, "Execução do PAINT")
    sheet.merge_cells("A3:S3")
    sheet["A3"] = f"{_filter_text(filters)} | {_paint_source_text(source)}"
    sheet["A3"].font = Font(size=9, italic=True, color="52647A")
    sheet["A3"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    cards = [
        (1, 4, "TOTAL DE AUDITORIAS", int(indicators.get("total_auditorias", 0) or 0), PAINT_KPI_COLORS[0]),
        (6, 9, "CONCLUÍDAS", int(indicators.get("concluidas", 0) or 0), PAINT_KPI_COLORS[1]),
        (11, 14, "EM ANDAMENTO", int(indicators.get("em_andamento", 0) or 0), PAINT_KPI_COLORS[2]),
        (16, 19, "NÃO INICIADAS", int(indicators.get("nao_iniciadas", 0) or 0), PAINT_KPI_COLORS[3]),
    ]
    for first_column, last_column, label, value, color in cards:
        _excel_kpi_card(sheet, 5, 8, first_column, last_column, label, value, color)

    _section_title(sheet, 11, "Visões da Execução do PAINT")
    _style_panel_range(sheet, 12, 38, 1, 9)
    _style_panel_range(sheet, 12, 38, 11, 19)
    sheet.merge_cells("A12:I12")
    sheet["A12"] = "Processos e Status"
    sheet["A12"].font = Font(size=11, bold=True, color="111111")
    sheet["A12"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("K12:S12")
    sheet["K12"] = "Gráfico por Status"
    sheet["K12"].font = Font(size=11, bold=True, color="111111")
    sheet["K12"].alignment = Alignment(horizontal="center")

    process_image = XLImage(_paint_process_chart_image(indicators))
    process_image.width = int(15.5 * 37.8)
    process_image.height = int(11.4 * 37.8)
    sheet.add_image(process_image, "A14")
    distribution_image = XLImage(_paint_distribution_chart_image(indicators))
    distribution_image.width = int(15.5 * 37.8)
    distribution_image.height = int(11.4 * 37.8)
    sheet.add_image(distribution_image, "K14")

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.print_area = "A1:S38"

    _write_paint_analytical_sheet(workbook, dataframe, filters)
    workbook.save(filename)
    return filename


def _carteira_status_chart_image(indicators):
    distribution = indicators.get("distribuicao_status", {})
    labels = list(CARTEIRA_STATUS_COLORS.keys())
    values = [
        int(distribution.get(label, {}).get("quantidade", 0) or 0)
        for label in labels
    ]
    return _pie_chart_image(
        labels,
        values,
        "Situação da Carteira",
        CARTEIRA_STATUS_COLORS,
        show_value_with_percent=True,
    )


def _carteira_year_chart_image(indicators):
    distribution = indicators.get("distribuicao_paint", {})
    labels = list(distribution.keys())
    values = [
        int(distribution.get(label, {}).get("quantidade", 0) or 0)
        for label in labels
    ]
    palette = {
        label: CARTEIRA_YEAR_COLORS[index % len(CARTEIRA_YEAR_COLORS)]
        for index, label in enumerate(labels)
    }
    return _pie_chart_image(
        labels,
        values,
        "Distribuição por PAINT",
        palette,
        show_value_with_percent=True,
    )


def _carteira_kpi_cards_pdf(indicators):
    cards = [
        ("CARTEIRA - AUDITORIAS", int(indicators.get("total_auditorias", 0) or 0)),
        (
            "AUDITORIAS COM PLANOS EM MONITORAMENTO",
            int(indicators.get("planos_monitoramento", 0) or 0),
        ),
        ("AUDITORIAS ENCERRADAS", int(indicators.get("encerradas", 0) or 0)),
        (
            "AGUARDANDO PLANO DE AÇÃO",
            int(indicators.get("aguardando_plano", 0) or 0),
        ),
    ]
    contents = [
        [
            _paragraph(label, size=8.2, bold=True, align=1),
            Spacer(1, 4),
            _paragraph(value, size=24, bold=True, align=1),
        ]
        for label, value in cards
    ]
    table = Table(
        [contents],
        colWidths=[95.5 * mm] * 4,
        rowHeights=[30 * mm],
        hAlign="CENTER",
    )
    commands = [
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#B8C7D9")),
        ("INNERGRID", (0, 0), (-1, -1), 6, colors.white),
        ("ROUNDEDCORNERS", [9]),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]
    for column, color in enumerate(CARTEIRA_KPI_COLORS):
        commands.append(("BACKGROUND", (column, 0), (column, 0), colors.HexColor(color)))
    table.setStyle(TableStyle(commands))
    return table


def _carteira_source_text(source):
    return (
        f"Base: {source.get('arquivo', 'carteira_auditoria.xlsx')}"
        f" | Período: {source.get('periodo', 'Não informado')}"
    )


def _carteira_pdf_table(dataframe, table_key):
    styles = getSampleStyleSheet()
    header_style = ParagraphStyle(
        f"CarteiraHeader{table_key}",
        parent=styles["BodyText"],
        fontName="Helvetica-Bold",
        fontSize=7,
        leading=8,
        textColor=colors.white,
        alignment=1,
    )
    cell_style = ParagraphStyle(
        f"CarteiraCell{table_key}",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=6.4,
        leading=7.6,
        textColor=colors.HexColor("#111111"),
    )
    columns = dataframe.columns.tolist()
    data = [[Paragraph(escape(str(column)), header_style) for column in columns]]
    for _, row in dataframe.iterrows():
        values = []
        for column in columns:
            value = row[column]
            if column == "% Implementação" and not pd.isna(value):
                text_value = f"{float(value):.1f}%".replace(".", ",")
            else:
                text_value = _safe_text(value)
            values.append(
                Paragraph(
                    escape(text_value).replace("\n", "<br/>"),
                    cell_style,
                )
            )
        data.append(values)

    if table_key == "aguardando":
        width_map = {
            "Auditorias": 94 * mm,
            "PAINT": 20 * mm,
            "Data da solicitação plano ação": 50 * mm,
            "Data limite para entrega do plano de ação*": 60 * mm,
            "Status": 36 * mm,
            "Data da entrega do plano de ação": 50 * mm,
            "Dias": 20 * mm,
        }
    else:
        width_map = {
            "Auditoria": 115 * mm,
            "PAINT": 24 * mm,
            "Status": 92 * mm,
            "Total Planos de Ação": 40 * mm,
            "% Implementação": 48 * mm,
            "Planos Abertos": 38 * mm,
        }
    widths = [width_map.get(column, 45 * mm) for column in columns]
    table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), BLUE),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.35, THIN_GREY),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT_BLUE]),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    if table_key == "aguardando" and "Status" in columns:
        status_column = columns.index("Status")
        status_colors = {"Vencido": "#FF3333"}
        for row_index, status in enumerate(dataframe["Status"].tolist(), start=1):
            background = status_colors.get(str(status))
            if background:
                commands.extend(
                    [
                        (
                            "BACKGROUND",
                            (status_column, row_index),
                            (status_column, row_index),
                            colors.HexColor(background),
                        ),
                        (
                            "FONTNAME",
                            (status_column, row_index),
                            (status_column, row_index),
                            "Helvetica-Bold",
                        ),
                    ]
                )
    table.setStyle(TableStyle(commands))
    return table


def export_carteira_pdf(indicators, carteira, aguardando, filters, source):
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    filename = EXPORT_DIR / f"carteira_auditoria_{_timestamp()}.pdf"
    document = _base_doc(filename)
    elements = []

    elements.extend(
        _pdf_header(
            "Indicadores e gráficos",
            filters,
            title="Carteira de Auditoria",
        )
    )
    elements.extend(
        [
            Paragraph(escape(_carteira_source_text(source)), getSampleStyleSheet()["Normal"]),
            Spacer(1, 3 * mm),
            _carteira_kpi_cards_pdf(indicators),
            Spacer(1, 6 * mm),
        ]
    )
    charts = Table(
        [
            [
                _panel_pdf(
                    "Situação da Carteira",
                    Image(_carteira_status_chart_image(indicators), width=182 * mm, height=110 * mm),
                    194 * mm,
                ),
                _panel_pdf(
                    "Distribuição por PAINT",
                    Image(_carteira_year_chart_image(indicators), width=182 * mm, height=110 * mm),
                    194 * mm,
                ),
            ]
        ],
        colWidths=[196 * mm, 196 * mm],
    )
    charts.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 2),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    elements.append(charts)

    elements.append(PageBreak())
    elements.extend(
        _pdf_header(
            "Tabelas Analíticas - Total da Carteira",
            filters,
            title="Carteira de Auditoria",
        )
    )
    elements.extend(
        [
            _paragraph("Carteira", size=12, bold=True),
            Spacer(1, 2 * mm),
            _carteira_pdf_table(carteira, "carteira"),
            Spacer(1, 4 * mm),
            _paragraph("Aguard. PA", size=12, bold=True),
            Spacer(1, 2 * mm),
            _carteira_pdf_table(aguardando, "aguardando"),
            Spacer(1, 2 * mm),
            _paragraph(
                "* Prazo limite é de 15 dias úteis a partir da notificação.",
                size=7,
            ),
        ]
    )

    document.build(
        elements,
        onFirstPage=_pdf_page_footer,
        onLaterPages=_pdf_page_footer,
    )
    return filename


def _write_carteira_analytical_sheet(workbook, dataframe, filters, table_key):
    sheet_name = "Aguard. PA" if table_key == "aguardando" else "Carteira"
    title = (
        "Tabela Analítica - Aguardando Plano de Ação"
        if table_key == "aguardando"
        else "Tabela Analítica - Total da Carteira"
    )
    sheet = workbook.create_sheet(sheet_name)
    sheet.sheet_view.showGridLines = False
    last_column = get_column_letter(max(1, len(dataframe.columns)))
    sheet.merge_cells(f"A1:{last_column}1")
    sheet["A1"] = title
    sheet["A1"].fill = _excel_fill(BLUE_HEX)
    sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet["A2"] = _filter_text(filters)
    sheet["A2"].font = Font(size=9, italic=True, color="52647A")
    sheet.row_dimensions[1].height = 28

    header_row = 4
    thin = Side(style="thin", color="9BAAC2")
    for column_index, column in enumerate(dataframe.columns, start=1):
        cell = sheet.cell(header_row, column_index, column)
        cell.fill = _excel_fill(BLUE_HEX)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    status_colors = {"Vencido": "#FF3333"}
    for row_index, (_, row) in enumerate(dataframe.iterrows(), start=header_row + 1):
        for column_index, column in enumerate(dataframe.columns, start=1):
            value = row[column]
            if pd.isna(value):
                value = ""
            elif hasattr(value, "item") and not isinstance(value, pd.Timestamp):
                value = value.item()
            cell = sheet.cell(row_index, column_index, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.fill = _excel_fill("#FFFFFF" if row_index % 2 else LIGHT_BLUE_HEX)
            if isinstance(value, (datetime, pd.Timestamp)):
                cell.number_format = "dd/mm/yyyy"
            if column == "% Implementação" and value != "":
                cell.number_format = '0.0"%"'
        if table_key == "aguardando" and "Status" in dataframe.columns:
            status_column = dataframe.columns.get_loc("Status") + 1
            status = str(row["Status"])
            background = status_colors.get(status)
            if background:
                status_cell = sheet.cell(row_index, status_column)
                status_cell.fill = _excel_fill(background)
                status_cell.font = Font(bold=True, color=_excel_font_color(background))

    widths = {
        "Auditoria": 52,
        "Auditorias": 46,
        "PAINT": 12,
        "Status": 34,
        "Total Planos de Ação": 22,
        "% Implementação": 20,
        "Planos Abertos": 18,
        "Data da solicitação plano ação": 24,
        "Data limite para entrega do plano de ação*": 31,
        "Data da entrega do plano de ação": 27,
        "Dias": 10,
    }
    for column_index, column in enumerate(dataframe.columns, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = widths.get(column, 24)
    last_row = header_row + len(dataframe)
    if table_key == "aguardando":
        note_row = last_row + 2
        sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=len(dataframe.columns))
        sheet.cell(note_row, 1, "* Prazo limite é de 15 dias úteis a partir da notificação.")
        sheet.cell(note_row, 1).font = Font(size=9, italic=True, color="52647A")
        last_row = note_row
    sheet.auto_filter.ref = f"A{header_row}:{last_column}{header_row + len(dataframe)}"
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:4"
    sheet.print_area = f"A1:{last_column}{last_row}"


def export_carteira_table_pdf(dataframe, filters, table_key):
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    filename = EXPORT_DIR / f"tabela_carteira_{table_key}_{_timestamp()}.pdf"
    document = _base_doc(filename)
    section = (
        "Aguardando Plano de Ação"
        if table_key == "aguardando"
        else "Total da Carteira"
    )
    elements = []
    elements.extend(
        _pdf_header(
            f"Tabela Analítica - {section}",
            filters,
            title="Carteira de Auditoria",
        )
    )
    elements.append(_carteira_pdf_table(dataframe, table_key))
    if table_key == "aguardando":
        elements.extend(
            [
                Spacer(1, 2 * mm),
                _paragraph(
                    "* Prazo limite é de 15 dias úteis a partir da notificação.",
                    size=7,
                ),
            ]
        )
    document.build(
        elements,
        onFirstPage=_pdf_page_footer,
        onLaterPages=_pdf_page_footer,
    )
    return filename


def export_carteira_table_excel(dataframe, filters, table_key):
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    filename = EXPORT_DIR / f"tabela_carteira_{table_key}_{_timestamp()}.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_carteira_analytical_sheet(workbook, dataframe, filters, table_key)
    workbook.save(filename)
    return filename


def export_carteira_excel(indicators, carteira, aguardando, filters, source):
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    filename = EXPORT_DIR / f"carteira_auditoria_{_timestamp()}.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Carteira de Auditoria"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"

    for column in range(1, 20):
        sheet.column_dimensions[get_column_letter(column)].width = 11.5
    for row in range(1, 75):
        sheet.row_dimensions[row].height = 20
    _paint_range(sheet, 1, 75, 1, 19, _excel_fill(PAGE_BACKGROUND_HEX))

    _merge_title(sheet, 1, "Carteira de Auditoria")
    sheet.merge_cells("A3:S3")
    sheet["A3"] = f"{_filter_text(filters)} | {_carteira_source_text(source)}"
    sheet["A3"].font = Font(size=9, italic=True, color="52647A")
    sheet["A3"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    cards = [
        (1, 4, "CARTEIRA - AUDITORIAS", int(indicators.get("total_auditorias", 0) or 0), CARTEIRA_KPI_COLORS[0]),
        (6, 9, "AUDITORIAS COM PLANOS EM MONITORAMENTO", int(indicators.get("planos_monitoramento", 0) or 0), CARTEIRA_KPI_COLORS[1]),
        (11, 14, "AUDITORIAS ENCERRADAS", int(indicators.get("encerradas", 0) or 0), CARTEIRA_KPI_COLORS[2]),
        (16, 19, "AGUARDANDO PLANO DE AÇÃO", int(indicators.get("aguardando_plano", 0) or 0), CARTEIRA_KPI_COLORS[3]),
    ]
    for first_column, last_column, label, value, color in cards:
        _excel_kpi_card(sheet, 5, 8, first_column, last_column, label, value, color)

    _section_title(sheet, 11, "Visões da Carteira de Auditoria")
    _style_panel_range(sheet, 12, 38, 1, 9)
    _style_panel_range(sheet, 12, 38, 11, 19)
    sheet.merge_cells("A12:I12")
    sheet["A12"] = "Situação da Carteira"
    sheet["A12"].font = Font(size=11, bold=True, color="111111")
    sheet["A12"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("K12:S12")
    sheet["K12"] = "Distribuição por PAINT"
    sheet["K12"].font = Font(size=11, bold=True, color="111111")
    sheet["K12"].alignment = Alignment(horizontal="center")

    status_image = XLImage(_carteira_status_chart_image(indicators))
    status_image.width = int(15.5 * 37.8)
    status_image.height = int(11.4 * 37.8)
    sheet.add_image(status_image, "A14")
    year_image = XLImage(_carteira_year_chart_image(indicators))
    year_image.width = int(15.5 * 37.8)
    year_image.height = int(11.4 * 37.8)
    sheet.add_image(year_image, "K14")

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.print_area = "A1:S38"

    _write_carteira_analytical_sheet(workbook, carteira, filters, "carteira")
    _write_carteira_analytical_sheet(workbook, aguardando, filters, "aguardando")
    workbook.save(filename)
    return filename
